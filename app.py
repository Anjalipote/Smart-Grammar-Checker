import html
import re
import difflib
import json
import io
import os
import base64
import random
from typing import List, Dict, Any, Tuple
import streamlit as st
import language_tool_python
import pypdf
import docx

# ---------------------------------------------------------------------------
# PAGE CONFIGURATION
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title="Smart Grammar Checker • AI Learning Studio",
    page_icon="logo.png" if os.path.exists("logo.png") else "✨",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ---------------------------------------------------------------------------
# VIBRANT LIGHT THEME WITH FAINT PASTEL GRADIENT BACKGROUND
# ---------------------------------------------------------------------------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;600&display=swap');

/* Global Styles - Faint Pastel Mesh Background */
html, body, [class*="css"], .stApp {
    font-family: 'Plus Jakarta Sans', -apple-system, BlinkMacSystemFont, sans-serif;
    background-color: #f8fafc !important;
    background-image: 
        radial-gradient(circle at 10% 12%, rgba(224, 231, 255, 0.8) 0%, transparent 48%),
        radial-gradient(circle at 90% 16%, rgba(209, 250, 229, 0.7) 0%, transparent 45%),
        radial-gradient(circle at 50% 80%, rgba(254, 226, 226, 0.55) 0%, transparent 50%),
        radial-gradient(circle at 82% 85%, rgba(254, 243, 199, 0.55) 0%, transparent 45%),
        radial-gradient(circle at 18% 78%, rgba(243, 232, 255, 0.6) 0%, transparent 48%) !important;
    color: #0f172a !important;
    letter-spacing: -0.01em;
}

/* Hide Default Streamlit Chrome */
header[data-testid="stHeader"] { visibility: hidden; height: 0px !important; }
#MainMenu { visibility: hidden; }
footer { visibility: hidden; }
.block-container {
    padding-top: 1.2rem !important;
    padding-bottom: 3rem !important;
    max-width: 1280px !important;
}

/* Light Frosted Hero Banner */
.hero-navbar {
    display: flex;
    align-items: center;
    justify-content: space-between;
    padding: 16px 26px;
    background: rgba(255, 255, 255, 0.88);
    backdrop-filter: blur(20px);
    -webkit-backdrop-filter: blur(20px);
    border: 1px solid rgba(226, 232, 240, 0.9);
    border-radius: 18px;
    margin-bottom: 22px;
    box-shadow: 0 10px 30px -5px rgba(99, 102, 241, 0.08), 0 2px 6px rgba(0, 0, 0, 0.02);
}
.hero-brand {
    display: flex;
    align-items: center;
    gap: 16px;
}
.hero-logo-img {
    width: 52px;
    height: 52px;
    border-radius: 14px;
    object-fit: cover;
    background: #ffffff;
    padding: 2px;
    box-shadow: 0 4px 14px rgba(0, 0, 0, 0.08);
    border: 1px solid rgba(226, 232, 240, 0.8);
}
.hero-title {
    font-size: 1.55rem;
    font-weight: 800;
    background: linear-gradient(135deg, #0f172a 0%, #312e81 60%, #0369a1 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    letter-spacing: -0.025em;
    margin: 0;
    line-height: 1.2;
}
.hero-tagline {
    font-size: 0.82rem;
    color: #475569;
    font-weight: 500;
    margin-top: 3px;
}
.hero-pills {
    display: flex;
    align-items: center;
    gap: 10px;
}
.pill-online {
    display: inline-flex;
    align-items: center;
    gap: 7px;
    padding: 6px 14px;
    background: #ecfdf5;
    border: 1px solid #a7f3d0;
    border-radius: 30px;
    font-size: 0.76rem;
    font-weight: 700;
    color: #047857;
    box-shadow: 0 2px 8px rgba(16, 185, 129, 0.1);
}
.pill-dot {
    width: 8px;
    height: 8px;
    background: #10b981;
    border-radius: 50%;
    box-shadow: 0 0 8px #10b981;
}

/* Learning Preset Cards */
.preset-grid {
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    gap: 12px;
    margin-bottom: 20px;
}
.preset-card {
    background: rgba(255, 255, 255, 0.85);
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 12px 14px;
    cursor: pointer;
    transition: all 0.25s cubic-bezier(0.4, 0, 0.2, 1);
    box-shadow: 0 4px 12px rgba(0, 0, 0, 0.03);
}
.preset-card:hover {
    transform: translateY(-2px);
    border-color: #6366f1;
    box-shadow: 0 8px 24px rgba(99, 102, 241, 0.12);
}
.preset-badge {
    display: inline-block;
    font-size: 0.68rem;
    font-weight: 700;
    text-transform: uppercase;
    letter-spacing: 0.06em;
    padding: 2px 7px;
    border-radius: 5px;
    margin-bottom: 6px;
}
.pb-agr { background: #fef3c7; color: #b45309; border: 1px solid #fde68a; }
.pb-hom { background: #e0f2fe; color: #0284c7; border: 1px solid #bae6fd; }
.pb-mod { background: #f3e8ff; color: #7e22ce; border: 1px solid #e9d5ff; }
.pb-ton { background: #fce7f3; color: #be185d; border: 1px solid #fbcfe8; }
.preset-desc {
    font-size: 0.78rem;
    color: #475569;
    line-height: 1.35;
    font-weight: 500;
}

/* Light Dashboard Cards with Faint Glow */
.glow-card {
    background: rgba(255, 255, 255, 0.9);
    backdrop-filter: blur(20px);
    border: 1px solid rgba(226, 232, 240, 0.9);
    border-radius: 16px;
    padding: 20px 24px;
    margin-bottom: 20px;
    box-shadow: 0 10px 30px -5px rgba(0, 0, 0, 0.04), 0 2px 6px rgba(0, 0, 0, 0.02);
    color: #0f172a;
}
.glow-card h3, .glow-card h4 {
    color: #0f172a !important;
    font-weight: 700;
}

/* Gamified Metrics Ribbon with Faint Pastel Tints */
.metrics-ribbon {
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(130px, 1fr));
    gap: 12px;
    margin-bottom: 22px;
}
.metric-tile {
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 14px 12px;
    text-align: center;
    position: relative;
    box-shadow: 0 4px 14px rgba(0, 0, 0, 0.03);
    transition: all 0.2s ease;
}
.metric-tile:hover {
    transform: translateY(-2px);
    box-shadow: 0 8px 24px rgba(0, 0, 0, 0.06);
}
.tile-val {
    font-size: 1.6rem;
    font-weight: 800;
    line-height: 1.1;
    margin-bottom: 4px;
}
.tile-val-emerald { color: #059669; }
.tile-val-amber { color: #d97706; }
.tile-val-blue { color: #0284c7; }
.tile-val-purple { color: #7c3aed; }
.tile-val-rose { color: #e11d48; }
.tile-lbl {
    font-size: 0.72rem;
    color: #64748b;
    text-transform: uppercase;
    font-weight: 700;
    letter-spacing: 0.05em;
}

/* Category Filter Tags with Soft Pastels */
.cat-tag {
    display: inline-block;
    padding: 3px 10px;
    border-radius: 6px;
    font-size: 0.72rem;
    font-weight: 700;
    letter-spacing: 0.04em;
    text-transform: uppercase;
}
.ct-agr { background: #fef3c7; color: #92400e; border: 1px solid #fde68a; }
.ct-spl { background: #ffe4e6; color: #9f1239; border: 1px solid #fecdd3; }
.ct-pnc { background: #e0f2fe; color: #0369a1; border: 1px solid #bae6fd; }
.ct-wrd { background: #ccfbf1; color: #0f766e; border: 1px solid #99f6e4; }
.ct-tns { background: #ede9fe; color: #5b21b6; border: 1px solid #ddd6fe; }
.ct-stl { background: #e0e7ff; color: #3730a3; border: 1px solid #c7d2fe; }
.ct-oth { background: #f1f5f9; color: #475569; border: 1px solid #cbd5e1; }

/* Interactive Learning Flashcard Diagnostics */
.tutor-card {
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 14px;
    padding: 18px 22px;
    margin-bottom: 16px;
    box-shadow: 0 4px 16px rgba(0, 0, 0, 0.04);
    transition: all 0.2s ease;
}
.tutor-card:hover {
    border-color: #6366f1;
    box-shadow: 0 8px 24px rgba(99, 102, 241, 0.09);
}
.tutor-header {
    display: flex;
    align-items: center;
    justify-content: space-between;
    margin-bottom: 12px;
}
.token-compare {
    display: flex;
    align-items: center;
    gap: 8px;
    font-size: 1rem;
    font-weight: 700;
}
.tok-wrong {
    background: #ffe4e6;
    color: #be123c;
    text-decoration: line-through;
    padding: 3px 9px;
    border-radius: 6px;
    border: 1px solid #fecdd3;
}
.tok-arrow {
    color: #94a3b8;
    font-size: 0.9rem;
}
.tok-right {
    background: #dcfce7;
    color: #15803d;
    padding: 3px 9px;
    border-radius: 6px;
    border: 1px solid #bbf7d0;
    box-shadow: 0 2px 6px rgba(22, 163, 74, 0.1);
}
.tutor-block {
    border-radius: 10px;
    padding: 12px 16px;
    font-size: 0.9rem;
    line-height: 1.55;
    margin-top: 10px;
}
.tb-wrong {
    background: #fff1f2;
    border-left: 4px solid #f43f5e;
    color: #881337;
}
.tb-right {
    background: #f0fdf4;
    border-left: 4px solid #10b981;
    color: #14532d;
}
.tb-protip {
    background: #eef2ff;
    border-left: 4px solid #6366f1;
    color: #312e81;
    margin-top: 8px;
    font-size: 0.86rem;
}

/* In-text Error Highlight */
.hl-bug {
    background: #fee2e2;
    border-bottom: 2px solid #ef4444;
    color: #991b1b;
    border-radius: 3px;
    padding: 1px 4px;
    cursor: help;
}

/* Diff Tags */
.diff-del {
    background: #ffe4e6;
    color: #be123c;
    text-decoration: line-through;
    padding: 2px 6px;
    border-radius: 4px;
}
.diff-add {
    background: #dcfce7;
    color: #15803d;
    font-weight: 700;
    padding: 2px 6px;
    border-radius: 4px;
}

/* Base Typography & Heading Contrast */
h1, h2, h3, h4, h5, h6, p, span, label, div {
    color: #0f172a;
}
.stMarkdown, .stMarkdown p, .stMarkdown span {
    color: #0f172a !important;
}

/* Textarea Styling - Clean Crisp White with High Contrast */
textarea,
.stTextArea textarea {
    background-color: #ffffff !important;
    color: #0f172a !important;
    -webkit-text-fill-color: #0f172a !important;
    border: 1.5px solid #cbd5e1 !important;
    border-radius: 14px !important;
    font-size: 1.02rem !important;
    font-weight: 500 !important;
    line-height: 1.65 !important;
    padding: 16px 18px !important;
    box-shadow: 0 2px 8px rgba(0, 0, 0, 0.03) !important;
    transition: all 0.2s ease !important;
}
textarea:focus,
.stTextArea textarea:focus {
    border-color: #6366f1 !important;
    box-shadow: 0 0 0 3px rgba(99, 102, 241, 0.18) !important;
}
textarea:disabled,
.stTextArea textarea:disabled {
    background-color: #f8fafc !important;
    color: #1e293b !important;
    -webkit-text-fill-color: #1e293b !important;
    opacity: 1 !important;
    border: 1.5px solid #cbd5e1 !important;
}

/* Explicit Placeholder Visibility */
textarea::placeholder,
.stTextArea textarea::placeholder,
textarea::-webkit-input-placeholder,
.stTextArea textarea::-webkit-input-placeholder,
textarea::-moz-placeholder,
.stTextArea textarea::-moz-placeholder {
    color: #64748b !important;
    -webkit-text-fill-color: #64748b !important;
    opacity: 1 !important;
    font-weight: 500 !important;
    font-size: 0.95rem !important;
}
input::placeholder,
.stTextInput input::placeholder,
input::-webkit-input-placeholder,
.stTextInput input::-webkit-input-placeholder,
input::-moz-placeholder,
.stTextInput input::-moz-placeholder {
    color: #64748b !important;
    -webkit-text-fill-color: #64748b !important;
    opacity: 1 !important;
    font-weight: 500 !important;
    font-size: 0.92rem !important;
}

/* Streamlit Inputs in Light Theme */
input,
.stTextInput input {
    background-color: #ffffff !important;
    color: #0f172a !important;
    -webkit-text-fill-color: #0f172a !important;
    border: 1.5px solid #cbd5e1 !important;
    border-radius: 10px !important;
    font-size: 0.95rem !important;
    font-weight: 500 !important;
    padding: 10px 14px !important;
}
input:focus,
.stTextInput input:focus {
    border-color: #6366f1 !important;
    box-shadow: 0 0 0 3px rgba(99, 102, 241, 0.18) !important;
}
div[data-baseweb="input"] {
    background-color: #ffffff !important;
    border-color: #cbd5e1 !important;
    border-radius: 10px !important;
}

/* Selectbox BaseWeb Overrides - Clean White & Dark Icons */
div[data-baseweb="select"],
div[data-baseweb="select"] > div,
div[data-baseweb="select"] [role="combobox"],
div[data-baseweb="select"] > div > div {
    background-color: #ffffff !important;
    color: #0f172a !important;
    border-color: #cbd5e1 !important;
    border-radius: 10px !important;
}
div[data-baseweb="select"] * {
    background-color: transparent !important;
    color: #0f172a !important;
    -webkit-text-fill-color: #0f172a !important;
    font-weight: 600 !important;
}
div[data-baseweb="select"] svg,
div[data-baseweb="select"] path {
    fill: #334155 !important;
    stroke: #334155 !important;
}
div[data-baseweb="popover"],
div[data-baseweb="popover"] > div,
div[data-baseweb="menu"],
ul[role="listbox"] {
    background-color: #ffffff !important;
    border: 1px solid #cbd5e1 !important;
    border-radius: 12px !important;
    box-shadow: 0 10px 30px rgba(0, 0, 0, 0.1) !important;
}
li[role="option"] {
    background-color: #ffffff !important;
    color: #0f172a !important;
    padding: 10px 14px !important;
}
li[role="option"] * {
    color: #0f172a !important;
    -webkit-text-fill-color: #0f172a !important;
    font-weight: 600 !important;
}
li[role="option"]:hover,
li[role="option"][aria-selected="true"] {
    background-color: #e0e7ff !important;
}
li[role="option"]:hover *,
li[role="option"][aria-selected="true"] * {
    color: #3730a3 !important;
    -webkit-text-fill-color: #3730a3 !important;
    font-weight: 700 !important;
}

/* Custom Tabs Styling - High Visibility Slate & Indigo */
.stTabs [data-baseweb="tab-list"] {
    gap: 8px !important;
    background: rgba(255, 255, 255, 0.95) !important;
    border: 1.5px solid #cbd5e1 !important;
    border-radius: 14px !important;
    padding: 6px !important;
    box-shadow: 0 4px 14px rgba(0, 0, 0, 0.04) !important;
}
.stTabs [data-baseweb="tab"],
.stTabs button[role="tab"] {
    border-radius: 10px !important;
    padding: 10px 18px !important;
    background: transparent !important;
    border: none !important;
    transition: all 0.2s ease !important;
}
.stTabs [data-baseweb="tab"] *,
.stTabs button[role="tab"] *,
.stTabs [data-baseweb="tab"] p,
.stTabs [data-baseweb="tab"] span,
.stTabs [data-baseweb="tab"] div {
    color: #334155 !important;
    -webkit-text-fill-color: #334155 !important;
    font-weight: 700 !important;
    font-size: 0.94rem !important;
}
.stTabs [data-baseweb="tab"]:hover *,
.stTabs button[role="tab"]:hover * {
    color: #1e293b !important;
    -webkit-text-fill-color: #1e293b !important;
}
.stTabs [aria-selected="true"],
.stTabs button[role="tab"][aria-selected="true"] {
    background: #ffffff !important;
    border-bottom: 3px solid #4f46e5 !important;
    box-shadow: 0 4px 12px rgba(79, 70, 229, 0.12) !important;
}
.stTabs [aria-selected="true"] *,
.stTabs button[role="tab"][aria-selected="true"] *,
.stTabs [aria-selected="true"] p,
.stTabs [aria-selected="true"] span,
.stTabs [aria-selected="true"] div {
    color: #4338ca !important;
    -webkit-text-fill-color: #4338ca !important;
    font-weight: 800 !important;
}

/* Captions & Subtitles */
.stCaption,
.stCaption p,
.stCaption span,
[data-testid="stCaptionContainer"],
[data-testid="stCaptionContainer"] p {
    color: #334155 !important;
    -webkit-text-fill-color: #334155 !important;
    font-size: 0.85rem !important;
    font-weight: 600 !important;
    line-height: 1.45 !important;
}

/* Radio Buttons & Checkboxes */
.stRadio label,
.stRadio div,
.stRadio p,
.stRadio span,
.stCheckbox label,
.stCheckbox div,
.stCheckbox p,
.stCheckbox span {
    color: #0f172a !important;
    -webkit-text-fill-color: #0f172a !important;
    font-weight: 600 !important;
}

/* Expanders */
[data-testid="stExpander"] {
    background: #ffffff !important;
    border: 1.5px solid #cbd5e1 !important;
    border-radius: 12px !important;
    box-shadow: 0 2px 8px rgba(0, 0, 0, 0.02) !important;
}
[data-testid="stExpander"] summary,
[data-testid="stExpander"] summary * {
    color: #0f172a !important;
    -webkit-text-fill-color: #0f172a !important;
    font-weight: 700 !important;
}
[data-testid="stExpander"] [data-testid="stExpanderDetails"],
[data-testid="stExpander"] [data-testid="stExpanderDetails"] * {
    color: #1e293b !important;
}

/* File Uploader Container */
[data-testid="stFileUploader"] {
    background: #ffffff !important;
    border-radius: 12px !important;
}
[data-testid="stFileUploader"] * {
    color: #1e293b !important;
}
[data-testid="stFileUploaderDropzone"] {
    background: #f8fafc !important;
    border: 2px dashed #94a3b8 !important;
    border-radius: 12px !important;
}
[data-testid="stFileUploaderDropzone"] * {
    color: #334155 !important;
}

/* Buttons */
.stButton>button {
    border-radius: 10px !important;
    font-weight: 700 !important;
    font-size: 0.92rem !important;
    padding: 10px 20px !important;
    transition: all 0.2s ease !important;
    background: #ffffff !important;
    color: #1e293b !important;
    border: 1.5px solid #cbd5e1 !important;
    box-shadow: 0 2px 6px rgba(0, 0, 0, 0.04) !important;
}
.stButton>button * {
    color: #1e293b !important;
    -webkit-text-fill-color: #1e293b !important;
}
.stButton>button:hover {
    transform: translateY(-1px) !important;
    border-color: #6366f1 !important;
    color: #4338ca !important;
    background: #f8fafc !important;
    box-shadow: 0 4px 14px rgba(99, 102, 241, 0.15) !important;
}
.stButton>button:hover * {
    color: #4338ca !important;
    -webkit-text-fill-color: #4338ca !important;
}

/* Primary Button */
.stButton>button[kind="primary"],
.stButton>button[data-testid="baseButton-primary"] {
    background: linear-gradient(135deg, #4f46e5, #6366f1) !important;
    color: #ffffff !important;
    border: none !important;
    box-shadow: 0 4px 14px rgba(79, 70, 229, 0.3) !important;
}
.stButton>button[kind="primary"] *,
.stButton>button[data-testid="baseButton-primary"] * {
    color: #ffffff !important;
    -webkit-text-fill-color: #ffffff !important;
}
.stButton>button[kind="primary"]:hover {
    color: #ffffff !important;
    box-shadow: 0 6px 20px rgba(79, 70, 229, 0.45) !important;
}

/* Tone Meters */
.tone-meter-bar {
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 16px 18px;
    margin-bottom: 14px;
    box-shadow: 0 2px 8px rgba(0, 0, 0, 0.02);
}

/* Quiz Container with Faint Violet Tint */
.quiz-box {
    background: linear-gradient(135deg, #f5f3ff 0%, #fdf4ff 100%);
    border: 1px solid #ddd6fe;
    border-radius: 14px;
    padding: 22px;
    margin-bottom: 20px;
    color: #1e1b4b;
    box-shadow: 0 8px 25px rgba(124, 58, 237, 0.08);
}
.quiz-box h3, .quiz-box div, .quiz-box p {
    color: #1e1b4b !important;
}

/* Duolingo Gamification Experience Styles */
.duo-hud-container {
    display: flex;
    align-items: center;
    justify-content: space-between;
    background: rgba(255, 255, 255, 0.96);
    border: 1.5px solid #cbd5e1;
    border-radius: 18px;
    padding: 12px 24px;
    margin-bottom: 20px;
    box-shadow: 0 6px 20px rgba(0, 0, 0, 0.04);
}
.duo-stat {
    display: flex;
    align-items: center;
    gap: 8px;
    font-size: 1.05rem;
    font-weight: 800;
}
.duo-hearts { color: #ef4444; }
.duo-streak { color: #ea580c; }
.duo-xp { color: #059669; }
.duo-level-badge {
    background: #eef2ff;
    color: #4338ca;
    padding: 6px 14px;
    border-radius: 30px;
    font-size: 0.85rem;
    font-weight: 800;
    border: 1px solid #c7d2fe;
}
.duo-arena-card {
    background: #ffffff;
    border: 2px solid #e2e8f0;
    border-radius: 20px;
    padding: 26px 30px;
    box-shadow: 0 10px 30px rgba(0, 0, 0, 0.05);
    margin-bottom: 20px;
}
.duo-mascot-row {
    display: flex;
    align-items: center;
    gap: 16px;
    margin-bottom: 20px;
}
.duo-mascot-avatar {
    font-size: 2.5rem;
    line-height: 1;
    background: #ecfdf5;
    border: 2px solid #a7f3d0;
    border-radius: 50%;
    padding: 8px 10px;
    box-shadow: 0 4px 12px rgba(16, 185, 129, 0.15);
}
.duo-bubble {
    background: #f8fafc;
    border: 1.5px solid #cbd5e1;
    border-radius: 14px;
    padding: 12px 18px;
    font-size: 1.05rem;
    font-weight: 700;
    color: #1e293b;
    flex: 1;
}
.duo-feedback-correct {
    background: #f0fdf4;
    border: 2px solid #86efac;
    border-radius: 16px;
    padding: 18px 22px;
    margin-top: 18px;
    color: #14532d;
}
.duo-feedback-wrong {
    background: #fff1f2;
    border: 2px solid #fecdd3;
    border-radius: 16px;
    padding: 18px 22px;
    margin-top: 18px;
    color: #881337;
}
.badge-grid {
    display: grid;
    grid-template-columns: repeat(4, 1fr);
    gap: 12px;
    margin-top: 14px;
}
.badge-item {
    background: #ffffff;
    border: 1.5px solid #e2e8f0;
    border-radius: 12px;
    padding: 12px 8px;
    text-align: center;
    box-shadow: 0 2px 8px rgba(0,0,0,0.03);
}
.badge-item.unlocked {
    border-color: #f59e0b;
    background: #fffbeb;
    box-shadow: 0 4px 14px rgba(245, 158, 11, 0.18);
}

/* Footer */
.app-footer {
    text-align: center;
    padding-top: 35px;
    font-size: 0.85rem;
    font-weight: 500;
    color: #475569;
    border-top: 1px solid rgba(226, 232, 240, 0.9);
    margin-top: 45px;
}
</style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# STRICTLY ONLINE CLOUD ENGINE
# ---------------------------------------------------------------------------
@st.cache_resource(show_spinner=False)
def get_cloud_tool(dialect: str):
    """
    Cached LanguageTool instance strictly routed to the official Cloud API.
    Zero local Java dependency, lightweight, high accuracy.
    """
    try:
        return language_tool_python.LanguageTool(dialect, remote_server="https://api.languagetool.org")
    except Exception:
        return language_tool_python.LanguageToolPublicAPI(dialect)

# ---------------------------------------------------------------------------
# COMPREHENSIVE LINGUISTIC LEARNING ENGINE
# ---------------------------------------------------------------------------
RULES_TUTOR_DATA = {
    "HE_VERB_AGR": {
        "category": "Subject-Verb & Agreement",
        "tag": "ct-agr",
        "why_wrong": "Third-person singular subjects ('he', 'she', 'it', or singular nouns) require verbs ending in '-s' or '-es' in present simple tense. Using a plural or base verb breaks syntactic agreement.",
        "why_right": "Using '{rep}' correctly inflects the verb to match the singular third-person subject.",
        "protip": "💡 Grammar Rule: He/She/It ALWAYS takes an -s in present tense (e.g. He walks, She thinks, It runs)."
    },
    "THEIR_IS": {
        "category": "Word Choice & Homophones",
        "tag": "ct-wrd",
        "why_wrong": "'Their' is a possessive pronoun expressing ownership (e.g., 'their car'). It cannot introduce an existential clause asserting the presence or location of items.",
        "why_right": "Using '{rep}' uses the existential pronoun 'there' ('there is / there are') to declare existence.",
        "protip": "💡 Memory Trick: 'THERE' has 'HERE' inside it (both denote place/existence). 'THEIR' has 'HEIR' (related to people owning things)."
    },
    "EVERYDAY_EVERY_DAY": {
        "category": "Word Choice & Homophones",
        "tag": "ct-wrd",
        "why_wrong": "'Everyday' (single word) is an adjective describing common items ('everyday routine'). It cannot function as an adverbial phrase modifying an action verb.",
        "why_right": "Using 'every day' (two words) supplies the temporal adverbial phrase meaning 'each day'.",
        "protip": "💡 Quick Test: Can you substitute 'each day'? If yes, use TWO words: 'every day'."
    },
    "MODAL_OF": {
        "category": "Tense & Verb Form",
        "tag": "ct-tns",
        "why_wrong": "Modal verbs ('should', 'could', 'would', 'must') can NEVER be paired with the preposition 'of'. This error happens because 'should've' phonetically sounds like 'should of'.",
        "why_right": "Using '{rep}' combines the modal verb with auxiliary verb 'have' to form the standard past modal perfect.",
        "protip": "💡 Golden Rule: In English grammar, 'should of', 'could of', and 'would of' NEVER exist. Always write 'should have'."
    },
    "MANY_NN": {
        "category": "Subject-Verb & Agreement",
        "tag": "ct-agr",
        "why_wrong": "Plural quantifiers (e.g. 'a few', 'many', 'several') must be paired with plural countable nouns, not singular forms.",
        "why_right": "Using '{rep}' correctly pluralizes the noun to maintain plural grammatical agreement.",
        "protip": "💡 Rule: If you can count them and have more than one ('a few', 'many'), the noun must end with plural -s."
    },
    "EACH_EVERY_NNS": {
        "category": "Subject-Verb & Agreement",
        "tag": "ct-agr",
        "why_wrong": "The distributives 'each' and 'every' focus on individuals one by one, so they grammatically govern singular nouns and singular verbs.",
        "why_right": "Using '{rep}' restores singular agreement according to distributive quantifier grammar.",
        "protip": "💡 Tip: Think of 'each' as 'each single one' — it always takes a singular verb (e.g. Each of them IS ready)."
    },
    "MODAL_VERB_WORD_ORDER": {
        "category": "Word Order & Syntax",
        "tag": "ct-syn",
        "why_wrong": "Modal verbs (will, would, can, could, shall, should, may, might, must) must appear BEFORE the base form of the main verb in English. Placing a conjugated verb before the modal creates an ungrammatical word order.",
        "why_right": "Using '{rep}' restores the correct English word order: modal + base verb (e.g. 'will have', 'can be', 'should do').",
        "protip": "💡 Rule: Modal verbs always come FIRST, followed by the base form of the verb (e.g. 'I will have', not 'I have will')."
    },
    "ITS_IT_IS": {
        "category": "Word Choice & Homophones",
        "tag": "ct-wrd",
        "why_wrong": "'Its' is a possessive pronoun (like 'his' or 'her'), whereas 'it's' is a contraction of 'it is' or 'it has'.",
        "why_right": "Using '{rep}' restores the intended grammatical meaning (either possession or contraction).",
        "protip": "💡 Quick Test: Try saying 'it is'. If 'it is' fits, write 'it's' with an apostrophe. If not, write 'its'."
    },
    "AFFECT_EFFECT": {
        "category": "Word Choice & Homophones",
        "tag": "ct-wrd",
        "why_wrong": "'Affect' is typically a verb (an Action meaning to influence), while 'effect' is typically a noun (an End result).",
        "why_right": "Using '{rep}' uses the correct grammatical part of speech required in this sentence.",
        "protip": "💡 Remember RAVEN: Remember Affect is a Verb, Effect is a Noun."
    },
    "A_VS_AN": {
        "category": "Word Choice & Phonetics",
        "tag": "ct-wrd",
        "why_wrong": "'A' must precede consonant sounds, while 'an' must precede vowel sounds (A, E, I, O, U sounds).",
        "why_right": "Using '{rep}' aligns with the phonetic sound of the next word, ensuring smooth pronunciation.",
        "protip": "💡 Note: It's about SOUND, not spelling! E.g. 'an hour' (silent h = vowel sound), but 'a university' (yoo sound = consonant)."
    },
    "DOUBLE_NEGATIVE": {
        "category": "Subject-Verb & Syntax",
        "tag": "ct-agr",
        "why_wrong": "Using two negatives in one clause (e.g. 'don't have no errors') cancels out or creates ungrammatical clutter.",
        "why_right": "Using '{rep}' keeps a single clean negation, delivering sharp and clear executive communication.",
        "protip": "💡 Tip: One negative per clause is all standard English allows (e.g. 'We don't have any errors')."
    },
    "IT_IS_2": {
        "category": "Word Choice & Homophones",
        "tag": "ct-wrd",
        "why_wrong": "'It's' is a contraction of 'it is' or 'it has', whereas 'its' is the possessive pronoun.",
        "why_right": "Using '{rep}' provides the correct possessive form without an apostrophe.",
        "protip": "💡 Memory Trick: If you cannot replace it with 'it is', write 'its' without an apostrophe."
    },
    "COLLECTIVE_NOUN_VERB_AGREEMENT_VBP": {
        "category": "Subject-Verb & Agreement",
        "tag": "ct-agr",
        "why_wrong": "In standard formal English, singular collective nouns ('the group', 'the team', 'the committee') take singular verbs ('has', 'is').",
        "why_right": "Using '{rep}' pairs the singular collective noun with a singular verb form.",
        "protip": "💡 Tip: Focus on the head noun ('The GROUP has concluded...'), not the plural object inside the phrase ('researchers')."
    },
    "PERS_PRONOUN_AGREEMENT": {
        "category": "Subject-Verb & Agreement",
        "tag": "ct-agr",
        "why_wrong": "The third-person plural pronoun 'they' requires the plural past verb 'were', not the singular 'was'.",
        "why_right": "Using '{rep}' ensures proper plural agreement with pronoun 'they'.",
        "protip": "💡 Rule: 'They were' / 'We were' / 'You were'. Use 'was' exclusively with singular subjects (I, he, she, it)."
    },
    "NON3PRS_VERB": {
        "category": "Subject-Verb & Agreement",
        "tag": "ct-agr",
        "why_wrong": "Third-person singular verb forms ending in '-s' ('needs', 'runs') cannot follow the first-person plural subject 'we'.",
        "why_right": "Using '{rep}' selects the uninflected base verb matching plural subject 'we'.",
        "protip": "💡 Tip: 'We', 'They', 'You', and 'I' always take base present verbs ('We need', 'They want', 'I speak')."
    },
    "OXFORD_SPELLING_Z_NOT_S": {
        "category": "Dialect & Regional Spelling",
        "tag": "ct-spl",
        "why_wrong": "The selected regional dialect follows regional spelling standards for '-ise' vs '-ize' word endings.",
        "why_right": "Using '{rep}' adheres to standard regional spelling conventions for this locale.",
        "protip": "💡 Regional Tip: US English prefers '-ize' (organize) while British/Commonwealth English commonly accepts '-ise' (organise)."
    }
}

class CustomMatch:
    def __init__(self, offset: int, error_length: int, rule_id: str, message: str, replacements: List[str], category: str = "Subject-Verb & Agreement"):
        self.offset = offset
        self.error_length = error_length
        self.rule_id = rule_id
        self.message = message
        self.replacements = replacements
        self.category = category

def check_custom_rules(text: str, existing_matches: list) -> list:
    """
    Augments LanguageTool public matches with deterministic linguistic heuristics
    for rules frequently omitted on the public tier (e.g. distributive quantifier agreement,
    modal-verb word-order).
    """
    augmented = list(existing_matches)

    # -----------------------------------------------------------------------
    # Rule: Modal-Verb Word-Order Correction
    # Detects misplaced modal verbs: "has will" -> "will have", etc.
    # When a conjugated verb form appears BEFORE a modal, the correct English
    # word order is modal + base-form verb.
    # -----------------------------------------------------------------------
    VERB_TO_BASE = {
        "has": "have", "have": "have", "had": "have",
        "is": "be", "am": "be", "are": "be", "was": "be", "were": "be",
        "does": "do", "do": "do", "did": "do",
    }
    MODALS = {"will", "would", "shall", "should", "can", "could", "may", "might", "must"}

    modal_pattern = re.compile(
        r"\b(" + "|".join(re.escape(v) for v in VERB_TO_BASE) + r")\s+(" + "|".join(MODALS) + r")\b",
        re.IGNORECASE
    )
    for m in modal_pattern.finditer(text):
        verb_form = m.group(1)
        modal = m.group(2)
        span_offset = m.start()
        span_length = m.end() - m.start()
        base = VERB_TO_BASE.get(verb_form.lower(), verb_form.lower())
        # Preserve capitalisation of the first word in the replacement
        if verb_form[0].isupper():
            replacement = f"{modal.capitalize()} {base}"
        else:
            replacement = f"{modal.lower()} {base}"
        # Remove any existing matches that overlap this span to avoid conflicts
        augmented = [
            ex for ex in augmented
            if not (ex.offset < span_offset + span_length and ex.offset + ex.error_length > span_offset)
        ]
        c_match = CustomMatch(
            offset=span_offset,
            error_length=span_length,
            rule_id="MODAL_VERB_WORD_ORDER",
            message=f"Modal verb word-order: '{verb_form} {modal}' should be '{replacement}'. Modal verbs (will, can, should …) must come before the base verb.",
            replacements=[replacement],
            category="Word Order & Syntax"
        )
        augmented.append(c_match)

    # -----------------------------------------------------------------------
    # Rule: Distributive Quantifier Agreement
    # Check for: (each|every one|neither|either) of the ... were/are -> was/is
    # -----------------------------------------------------------------------
    pattern = re.compile(r"\b(each|every\s+one|neither|either)\s+of\s+(?:the|these|those)\s+[a-z]+\s+(were|are)\b", re.IGNORECASE)
    for m in pattern.finditer(text):
        verb = m.group(2)
        verb_offset = m.start(2)
        verb_len = len(verb)
        overlap = any(
            ex.offset <= verb_offset < (ex.offset + ex.error_length) or
            verb_offset <= ex.offset < (verb_offset + verb_len)
            for ex in augmented
        )
        if not overlap:
            replacement = "was" if verb.lower() == "were" else "is"
            if verb[0].isupper():
                replacement = replacement.capitalize()
            c_match = CustomMatch(
                offset=verb_offset,
                error_length=verb_len,
                rule_id="EACH_EVERY_NNS",
                message=f"Distributive quantifier agreement: '{m.group(1)}' takes a singular verb ('{replacement}').",
                replacements=[replacement],
                category="Subject-Verb & Agreement"
            )
            augmented.append(c_match)
    augmented.sort(key=lambda x: x.offset)
    return augmented

def is_whitelisted(token: str, whitelist: set) -> bool:
    if not whitelist:
        return False
    clean = token.strip(".,!?;:\"'()[]{} \t\n\r").lower()
    if clean in whitelist:
        return True
    tok_words = clean.split()
    if any(w in whitelist for w in tok_words):
        return True
    return False

def clean_txt(s: str) -> str:
    if not s:
        return ""
    return (s.replace("‘", "'")
            .replace("’", "'")
            .replace("“", '"')
            .replace("”", '"')
            .replace("`", "'")
            .replace("\ufffd", "'"))

def analyze_rule(match, token: str) -> Dict[str, Any]:
    rule_id = getattr(match, "rule_id", "") or ""
    cat_raw = (getattr(match, "category", "") or "").upper()
    raw_msg = clean_txt(getattr(match, "message", "") or "Grammar inconsistency detected.")
    replacements = getattr(match, "replacements", []) or []
    best_rep = replacements[0] if replacements else ""

    if rule_id in RULES_TUTOR_DATA:
        info = RULES_TUTOR_DATA[rule_id]
        return {
            "category": info["category"],
            "tag": info["tag"],
            "why_wrong": info["why_wrong"],
            "why_right": info["why_right"].format(rep=best_rep if best_rep else "the correction"),
            "protip": info["protip"],
            "rule_id": rule_id,
            "message": raw_msg,
            "best_replacement": best_rep
        }

    # Dynamic Classifier
    if any(k in rule_id or k in cat_raw for k in ["VERB_AGR", "AGREEMENT", "PLURAL", "SINGULAR", "SUBJECT"]):
        category = "Subject-Verb & Agreement"
        tag = "ct-agr"
        why_wrong = f"Subject-verb disagreement: The subject and verb do not match in number or person. {raw_msg}"
        why_right = f"Using '{best_rep}' restores number/person harmony between subject and verb."
        protip = "💡 Master Tip: Find the true subject of the clause and verify if it's singular or plural before picking the verb."
    elif any(k in rule_id or k in cat_raw for k in ["TYPO", "SPELL", "MORFOLOGIK", "MISSING"]):
        category = "Spelling & Typo"
        tag = "ct-spl"
        why_wrong = f"The word '{token}' is misspelled or not found in standard English dictionaries."
        why_right = f"'{best_rep}' is the standard, correctly spelled English word."
        protip = "💡 Master Tip: Double check silent letters and double consonants when spelling this word."
    elif any(k in rule_id or k in cat_raw for k in ["PUNCTUATION", "COMMA", "PERIOD", "APOSTROPHE", "COLON", "SPACE"]):
        category = "Punctuation & Mechanics"
        tag = "ct-pnc"
        why_wrong = f"Punctuation or spacing violation: {raw_msg}"
        why_right = f"Applying '{best_rep}' restores standard typographical clarity."
        protip = "💡 Master Tip: Read the sentence aloud; natural pauses help you place commas and periods accurately."
    elif any(k in rule_id or k in cat_raw for k in ["CONFUSED", "HOMOPHONE", "PREPOSITION"]):
        category = "Word Choice & Homophones"
        tag = "ct-wrd"
        why_wrong = f"Homophone or word choice error: {raw_msg}"
        why_right = f"Using '{best_rep}' supplies the correct word with the intended meaning."
        protip = "💡 Master Tip: Homophones sound identical but have different meanings and spellings."
    elif any(k in rule_id or k in cat_raw for k in ["TENSE", "MODAL", "IRREGULAR", "GERUND", "PARTICIPLE"]):
        category = "Tense & Verb Form"
        tag = "ct-tns"
        why_wrong = f"Verb form or tense inconsistency: {raw_msg}"
        why_right = f"Using '{best_rep}' puts the verb into the proper grammatical tense."
        protip = "💡 Master Tip: Keep verb tenses consistent throughout your paragraph unless expressing time shifts."
    else:
        category = "Style & Clarity"
        tag = "ct-stl"
        why_wrong = f"Stylistic or syntactic issue: {raw_msg}"
        why_right = f"Replacing with '{best_rep}' improves clarity and readability."
        protip = "💡 Master Tip: Aim for clear, direct, and concise sentence structures."

    return {
        "category": category,
        "tag": tag,
        "why_wrong": why_wrong,
        "why_right": why_right,
        "protip": protip,
        "rule_id": rule_id,
        "message": raw_msg,
        "best_replacement": best_rep
    }

# ---------------------------------------------------------------------------
# TONE & VOCABULARY POWER-UP ENGINE
# ---------------------------------------------------------------------------
CONTRACTIONS = {"can't", "won't", "don't", "didn't", "it's", "i'm", "you're", "they're", "we're", "shouldn't", "couldn't", "wouldn't", "isn't", "aren't", "wasn't", "gonna", "wanna"}
INFORMAL_WORDS = {"hey", "yeah", "cool", "stuff", "awesome", "dude", "nope", "yep", "tbh", "btw", "ok", "okay"}
FORMAL_CONNECTORS = {"furthermore", "consequently", "moreover", "therefore", "additionally", "accordingly", "nevertheless"}
HEDGE_WORDS = {"maybe", "perhaps", "possibly", "i think", "sort of", "kind of", "might"}
COURTESY_WORDS = {"please", "thank you", "appreciate", "kindly", "grateful", "cordially", "sincerely"}

POWER_VOCAB_MAP = {
    "very big": ["colossal", "immense", "substantial"],
    "very good": ["exceptional", "exemplary", "superb"],
    "very bad": ["detrimental", "deficient", "adverse"],
    "very important": ["paramount", "critical", "pivotal"],
    "very easy": ["effortless", "straightforward", "accessible"],
    "a lot of": ["an abundance of", "numerous", "substantial"],
    "help": ["facilitate", "assist", "expedite"],
    "make sure": ["ensure", "verify", "ascertain"],
    "start": ["commence", "initiate", "embark upon"],
    "show": ["demonstrate", "illustrate", "exemplify"]
}

def analyze_tone_and_vocab(text: str, target_tone: str = "Executive / Business 👔") -> Dict[str, Any]:
    if not text.strip():
        return {
            "persona": "Neutral & Clear",
            "formality": 50,
            "directness": 50,
            "politeness": 50,
            "summary": "Ready to evaluate tone and provide vocabulary power-ups upon text input.",
            "vocab_suggestions": [],
            "target_tone": target_tone,
            "target_alignment_score": 50,
            "target_alignment_status": "Ready",
            "target_coaching": [f"Enter text to analyze alignment with {target_tone}"]
        }

    words = re.findall(r"\b[A-Za-z']+\b", text.lower())
    n = max(1, len(words))

    c_found = [w for w in words if w in CONTRACTIONS]
    inf_found = [w for w in words if w in INFORMAL_WORDS]
    form_found = [w for w in words if w in FORMAL_CONNECTORS]
    avg_len = sum(len(w) for w in words) / n

    formality = max(10, min(95, int(65 - (len(c_found) * 5) - (len(inf_found) * 6) + (len(form_found) * 7) + (10 if avg_len > 5.1 else -10 if avg_len < 4.0 else 0))))
    hedge_found = [h for h in HEDGE_WORDS if h in text.lower()]
    passive_count = len(re.findall(r"\b(is|are|was|were|been|being)\s+[a-z]+ed\b", text.lower()))
    directness = max(15, min(95, int(80 - (len(hedge_found) * 8) - (passive_count * 6))))
    courtesy_found = [c for c in COURTESY_WORDS if c in text.lower()]
    politeness = min(95, 50 + (len(courtesy_found) * 15))

    if formality >= 75 and directness >= 65:
        persona = "Executive & Authoritative 👔"
        summary = "Confident, disciplined, and refined. Ideal for leadership memos, board pitches, and executive proposals."
    elif formality >= 70:
        persona = "Academic & Scholarly 🎓"
        summary = "Analytical, precise, and rigorous vocabulary. Suited for research publications and whitepapers."
    elif formality <= 45:
        persona = "Casual & Conversational ☕"
        summary = "Approachable, friendly, and lively. Great for team chats, blog posts, and creative stories."
    elif directness >= 80:
        persona = "Direct & Impactful ⚡"
        summary = "Laser-focused and decisive with high active voice. Ideal for urgent calls to action."
    else:
        persona = "Professional & Balanced 🌟"
        summary = "Balanced professional voice combining courtesy, clarity, and ease of reading."

    vocab_hits = []
    text_lower = text.lower()
    for weak_term, strong_opts in POWER_VOCAB_MAP.items():
        if weak_term in text_lower:
            vocab_hits.append({
                "weak": weak_term,
                "strong": ", ".join(strong_opts)
            })

    # Target Tone Strategic Alignment Calculation
    coaching = []
    if "Executive" in target_tone:
        score = int((formality * 0.45) + (directness * 0.45) + (politeness * 0.10))
        if inf_found:
            score -= len(inf_found) * 12
            coaching.append(f"⚠️ Replace informal greeting/words ({', '.join(set(inf_found))}) with professional executive language.")
        if c_found:
            score -= len(c_found) * 4
            coaching.append(f"💡 Expand contractions ({', '.join(set(c_found))}) for formal executive writing.")
        if hedge_found:
            score -= len(hedge_found) * 8
            coaching.append(f"📌 Replace hesitant words ({', '.join(set(hedge_found))}) with decisive executive statements.")
        for v in vocab_hits:
            coaching.append(f"💎 Upgrade '{v['weak']}' with executive power vocabulary: {v['strong']}.")
        if "don't have no" in text_lower or len(re.findall(r"\b(?:don't|doesn't|didn't|never|cannot)\b[^\.\?!]*\b(?:no|nobody|nothing|nowhere)\b", text_lower)) > 0:
            score -= 15
            coaching.append("⚠️ Remove double negatives to project boardroom authority and clarity.")
        if not coaching:
            coaching.append("✅ Outstanding executive tone! Clear, authoritative, and boardroom-ready.")

    elif "Academic" in target_tone:
        score = int((formality * 0.65) + (politeness * 0.20) + (15 if form_found else 0))
        if c_found:
            score -= len(c_found) * 8
            coaching.append(f"⚠️ Academic style strictly avoids contractions ({', '.join(set(c_found))}); spell words out completely.")
        if inf_found:
            score -= len(inf_found) * 15
            coaching.append(f"⚠️ Remove colloquial words ({', '.join(set(inf_found))}) in favor of objective academic vocabulary.")
        if not form_found:
            coaching.append("💡 Incorporate formal discourse transitions ('furthermore', 'consequently', 'moreover') to connect arguments logically.")
        if not coaching:
            coaching.append("✅ Scholarly tone verified! Formal, precise, and well-structured.")

    elif "Direct" in target_tone:
        score = int((directness * 0.8) + (formality * 0.2))
        if hedge_found:
            score -= len(hedge_found) * 14
            coaching.append(f"⚠️ Eliminate hedging ({', '.join(set(hedge_found))}) to maximize decisive impact.")
        if passive_count > 0:
            score -= passive_count * 10
            coaching.append("💡 Convert passive voice constructions into active, subject-driven actions.")
        if not coaching:
            coaching.append("✅ High-impact voice! Laser-focused, action-oriented, and decisive.")

    else:  # Casual / Friendly
        score = int(100 - (formality * 0.5) + (politeness * 0.3))
        if formality >= 75:
            score -= 20
            coaching.append("💡 The phrasing is quite formal. Introduce conversational ease, friendly greetings, or relatable phrasing.")
        if not courtesy_found and not inf_found:
            coaching.append("💡 Add warmth or a personal touch to create an approachable connection.")
        if not coaching:
            coaching.append("✅ Warm and conversational! Engaging and easy for readers to connect with.")

    score = max(10, min(100, score))
    status = "Aligned 🎯" if score >= 75 else "Moderate Fit ⚖️" if score >= 50 else "Needs Adjustment ⚠️"

    return {
        "persona": persona,
        "formality": formality,
        "directness": directness,
        "politeness": politeness,
        "summary": summary,
        "vocab_suggestions": vocab_hits,
        "target_tone": target_tone,
        "target_alignment_score": score,
        "target_alignment_status": status,
        "target_coaching": coaching
    }

def count_syllables(w: str) -> int:
    w = w.lower().strip()
    if len(w) <= 3:
        return 1
    w = re.sub(r'(?:[^laeiouy]|ed|es|e)$', '', w)
    w = re.sub(r'^y', '', w)
    matches = re.findall(r'[aeiouy]{1,2}', w)
    return max(1, len(matches))

def calculate_readability(text: str) -> Dict[str, Any]:
    words = re.findall(r'\b[A-Za-z]+\b', text)
    sentences = [s for s in re.split(r'[.!?]+', text) if s.strip()]

    n_words = len(words)
    n_sentences = max(1, len(sentences))

    if n_words == 0:
        return {"words": 0, "sentences": 0, "grade": 0.0, "read_time": 0.0, "level": "N/A"}

    total_syllables = sum(count_syllables(w) for w in words)
    asl = n_words / n_sentences
    asw = total_syllables / n_words

    reading_ease = max(0.0, min(100.0, 206.835 - (1.015 * asl) - (84.6 * asw)))
    grade = max(0.0, min(18.0, (0.39 * asl) + (11.8 * asw) - 15.59))

    if reading_ease >= 70:
        level = "Very Easy / Plain English"
    elif reading_ease >= 55:
        level = "Standard Business"
    else:
        level = "Advanced Academic"

    return {
        "words": n_words,
        "sentences": n_sentences,
        "grade": round(grade, 1),
        "read_time": round(n_words / 200, 1),
        "level": level
    }

def build_visual_diff(original: str, corrected: str) -> str:
    orig_words = original.split()
    corr_words = corrected.split()
    matcher = difflib.SequenceMatcher(None, orig_words, corr_words)
    diff_html = []

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == 'equal':
            diff_html.append(html.escape(" ".join(orig_words[i1:i2])))
        elif tag == 'delete':
            diff_html.append(f'<span class="diff-del">{html.escape(" ".join(orig_words[i1:i2]))}</span>')
        elif tag == 'insert':
            diff_html.append(f'<span class="diff-add">{html.escape(" ".join(corr_words[j1:j2]))}</span>')
        elif tag == 'replace':
            diff_html.append(
                f'<span class="diff-del">{html.escape(" ".join(orig_words[i1:i2]))}</span> '
                f'<span class="diff-add">{html.escape(" ".join(corr_words[j1:j2]))}</span>'
            )
    return " ".join(diff_html)

def build_annotated_preview(text: str, matches) -> str:
    if not matches:
        return html.escape(text)

    sorted_m = sorted(matches, key=lambda m: m.offset)
    out = []
    last_idx = 0

    for m in sorted_m:
        start = m.offset
        end = m.offset + m.error_length
        if start < last_idx:
            continue

        out.append(html.escape(text[last_idx:start]))
        err_token = html.escape(text[start:end])
        clean_msg = clean_txt(m.message)
        out.append(f'<span class="hl-bug" title="{html.escape(clean_msg)}">{err_token}</span>')
        last_idx = end

    out.append(html.escape(text[last_idx:]))
    return "".join(out).replace("\n", "<br>")

def parse_uploaded_file(uploaded_file) -> Tuple[str, Dict[str, Any], str]:
    fname = uploaded_file.name.lower()
    text = ""
    err = ""

    try:
        if fname.endswith(".txt") or fname.endswith(".md"):
            raw = uploaded_file.read()
            try:
                text = raw.decode("utf-8")
            except UnicodeDecodeError:
                text = raw.decode("latin-1", errors="replace")

        elif fname.endswith(".docx") or fname.endswith(".doc"):
            if fname.endswith(".doc"):
                err = "Legacy binary .doc files cannot be parsed directly. Please save or export your file as modern .docx or plain .txt before uploading."
            else:
                doc = docx.Document(uploaded_file)
                paras = [p.text for p in doc.paragraphs if p.text.strip()]
                text = "\n\n".join(paras)

        elif fname.endswith(".pdf"):
            reader = pypdf.PdfReader(uploaded_file)
            pages = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    pages.append(extracted.strip())
            text = "\n\n".join(pages)
            if not text.strip():
                err = "The PDF contains no extractable text (it might be a scanned image or protected PDF)."

        elif fname.endswith(".csv"):
            raw = uploaded_file.read()
            text = raw.decode("utf-8", errors="replace")

        else:
            raw = uploaded_file.read()
            text = raw.decode("utf-8", errors="replace")

    except Exception as e:
        err = f"Failed to extract document contents: {str(e)}"

    meta = {
        "filename": uploaded_file.name,
        "size_kb": round(len(uploaded_file.getvalue()) / 1024, 1),
        "word_count": len(re.findall(r"\b\w+\b", text))
    }
    return text, meta, err

def export_docx_file(original: str, corrected: str, diagnostics: List[Dict[str, Any]], dialect: str) -> io.BytesIO:
    doc = docx.Document()
    if os.path.exists("logo.png"):
        try:
            doc.add_picture("logo.png", width=docx.shared.Inches(1.2))
        except Exception:
            pass
    doc.add_heading("Smart Grammar Checker • Polished Document", level=0)
    p = doc.add_paragraph()
    p.add_run(f"Dialect: {dialect} | Engine: Online Cloud API | Issues Resolved: {len(diagnostics)}").italic = True

    doc.add_heading("Corrected Document", level=1)
    for para in corrected.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    doc.add_page_break()
    doc.add_heading("Diagnostic Audit Report", level=1)
    if not diagnostics:
        doc.add_paragraph("No grammatical or stylistic issues identified. Perfect execution!")
    else:
        for idx, d in enumerate(diagnostics, 1):
            p_diag = doc.add_paragraph()
            p_diag.add_run(f"{idx}. [{d['category']}] ").bold = True
            p_diag.add_run(f"'{d['token']}' ➔ '{d['best_replacement']}'\n").bold = True
            p_diag.add_run("• Why It's Wrong: ").bold = True
            p_diag.add_run(f"{d['why_wrong']}\n")
            p_diag.add_run("• Why It's Right: ").bold = True
            p_diag.add_run(f"{d['why_right']}\n")
            if d.get("protip"):
                p_diag.add_run(f"• Pro Tip: {d['protip']}\n")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

# ---------------------------------------------------------------------------
# DUOLINGO GRAMMAR QUEST DATA (5 PROGRESSIVE LEVELS • 25 QUESTIONS)
# ---------------------------------------------------------------------------
DUOLINGO_LEVELS = [
    {
        "level_id": 1,
        "title": "Level 1: Subject-Verb Agreement Apprentice",
        "description": "Master singular vs. plural subjects, quantifiers, and collective nouns.",
        "icon": "🌱",
        "badge": "Agreement Apprentice 🎖️",
        "questions": [
            {
                "q": "Neither the teacher nor the students _____ prepared for the surprise inspection.",
                "options": ["was", "were"],
                "correct": "were",
                "explanation": "With 'neither...nor', the verb agrees with the subject closer to it. 'Students' is plural, so use 'were'.",
                "tip": "💡 Memory Hack: Look at the word right before the verb! (Students ➔ were)"
            },
            {
                "q": "The group of researchers _____ discovered a breakthrough vaccine.",
                "options": ["has", "have"],
                "correct": "has",
                "explanation": "The true subject is 'The group' (a singular collective noun), not 'researchers'. Therefore, it takes singular 'has'.",
                "tip": "💡 Memory Hack: Ignore words in prepositional phrases like 'of researchers' when finding the subject!"
            },
            {
                "q": "Each of the new participants _____ provided with an identification badge.",
                "options": ["is", "are"],
                "correct": "is",
                "explanation": "'Each' is a distributive pronoun that always governs a singular verb.",
                "tip": "💡 Memory Hack: Think of 'each' as 'each single one'."
            },
            {
                "q": "Every member of the committee _____ voted on the new fiscal policy.",
                "options": ["has", "have"],
                "correct": "has",
                "explanation": "'Every' modifies 'member', which is singular. Thus, singular 'has' is required.",
                "tip": "💡 Memory Hack: 'Every' and 'Each' are always singular champions."
            },
            {
                "q": "A few _____ of advice were shared during the opening conference.",
                "options": ["pieces", "piece"],
                "correct": "pieces",
                "explanation": "Plural quantifier 'a few' requires a plural countable noun ('pieces').",
                "tip": "💡 Memory Hack: 'Advice' is uncountable, so count it with plural 'pieces'!"
            }
        ]
    },
    {
        "level_id": 2,
        "title": "Level 2: Tricky Homophones & Word Choice",
        "description": "Never confuse their/there, its/it's, and affect/effect again.",
        "icon": "🔀",
        "badge": "Homophone Hunter 🎯",
        "questions": [
            {
                "q": "How will the recent interest rate changes _____ our corporate investment plan?",
                "options": ["affect", "effect"],
                "correct": "affect",
                "explanation": "'Affect' is an Action verb (to influence). 'Effect' is a Noun (the result). Here an action verb is required.",
                "tip": "💡 Memory Hack: Remember RAVEN — Remember Affect is a Verb, Effect is a Noun!"
            },
            {
                "q": "The company celebrated _____ ten-year anniversary in New York.",
                "options": ["its", "it's"],
                "correct": "its",
                "explanation": "'Its' without apostrophe is the possessive pronoun. 'It's' is a contraction of 'it is'.",
                "tip": "💡 Memory Hack: Test with 'it is'. 'The company celebrated it is anniversary' makes no sense, so use 'its'!"
            },
            {
                "q": "_____ is an urgent message waiting for you at the front desk.",
                "options": ["There", "Their"],
                "correct": "There",
                "explanation": "'There' indicates place or existence ('there is / there are'). 'Their' is possessive ownership.",
                "tip": "💡 Memory Hack: 'THERE' contains 'HERE' — both denote place!"
            },
            {
                "q": "Our team holds an agile stand-up meeting _____ at 9:00 AM.",
                "options": ["every day", "everyday"],
                "correct": "every day",
                "explanation": "'Every day' (two words) is an adverbial time phrase meaning 'each day'. 'Everyday' (one word) is an adjective.",
                "tip": "💡 Memory Hack: Replace it with 'each day'. If 'each day' fits, write TWO words!"
            },
            {
                "q": "I would rather take the express train _____ drive through rush hour traffic.",
                "options": ["than", "then"],
                "correct": "than",
                "explanation": "'Than' is used for comparisons. 'Then' refers to time and chronology ('first this, then that').",
                "tip": "💡 Memory Hack: CompAre uses thAn. TimE uses thEn."
            }
        ]
    },
    {
        "level_id": 3,
        "title": "Level 3: Tense & Modal Auxiliary Perfection",
        "description": "Conquer modal verbs, past conditionals, and irregular past tenses.",
        "icon": "⏱️",
        "badge": "Time Bender ⏳",
        "questions": [
            {
                "q": "I should _____ verified the server credentials before deploying the release.",
                "options": ["have", "of"],
                "correct": "have",
                "explanation": "'Should of' NEVER exists in English! Modal auxiliary verbs always pair with 'have' ('should have').",
                "tip": "💡 Memory Hack: It sounds like 'of' when spoken quickly, but it is ALWAYS 'have'."
            },
            {
                "q": "By the time the keynote speaker arrived, the audience _____ seated.",
                "options": ["had been", "have been"],
                "correct": "had been",
                "explanation": "Actions that happened prior to another past event require past perfect tense ('had been').",
                "tip": "💡 Memory Hack: The earlier past event always takes 'had' + past participle."
            },
            {
                "q": "If he _____ studied the technical rubric, he would have scored higher.",
                "options": ["had", "would have"],
                "correct": "had",
                "explanation": "In third conditional 'if' clauses, use past perfect ('if he had studied'), never 'would have' inside the 'if' condition.",
                "tip": "💡 Memory Hack: No 'would' in the 'if' clause!"
            },
            {
                "q": "The data scientist has _____ the financial forecast twice.",
                "options": ["proven", "proved"],
                "correct": "proven",
                "explanation": "'Proven' is the preferred standard past participle in formal analytical writing.",
                "tip": "💡 Memory Hack: Has + proven (participle harmony)."
            },
            {
                "q": "She could not _____ expected such an overwhelming customer response.",
                "options": ["have", "of"],
                "correct": "have",
                "explanation": "'Could of' is ungrammatical; standard English strictly requires 'could have'.",
                "tip": "💡 Memory Hack: Modal + have (could have, would have, should have)."
            }
        ]
    },
    {
        "level_id": 4,
        "title": "Level 4: Executive Tone & Power Vocabulary",
        "description": "Eliminate double negatives, upgrade weak adjectives, and master active voice.",
        "icon": "👔",
        "badge": "Executive Maestro 👑",
        "questions": [
            {
                "q": "Select the professional alternative: 'We need to _____ that all security tests pass.'",
                "options": ["ensure", "make sure"],
                "correct": "ensure",
                "explanation": "'Ensure' is the concise, authoritative executive equivalent of the colloquial phrase 'make sure'.",
                "tip": "💡 Memory Hack: Executive writing prefers single precise verbs over two-word phrases."
            },
            {
                "q": "Identify the grammatically correct sentence:",
                "options": [
                    "We don't have any pending audit findings.",
                    "We don't have no pending audit findings."
                ],
                "correct": "We don't have any pending audit findings.",
                "explanation": "Double negatives ('don't have no') create ungrammatical confusion. Use a single negative with 'any'.",
                "tip": "💡 Memory Hack: In English grammar, one negation per clause is the gold standard."
            },
            {
                "q": "Select the highest impact power verb: 'The new dashboard will _____ our reporting workflow.'",
                "options": ["accelerate", "help do"],
                "correct": "accelerate",
                "explanation": "Strong action verbs ('accelerate', 'expedite') convey dynamism and executive leadership.",
                "tip": "💡 Memory Hack: Strong verbs make your achievements sound impactful."
            },
            {
                "q": "Which phrasing demonstrates active, authoritative voice?",
                "options": [
                    "The committee approved the strategic expansion.",
                    "The strategic expansion was approved by the committee."
                ],
                "correct": "The committee approved the strategic expansion.",
                "explanation": "Active voice puts the actor ('The committee') before the verb, delivering clarity and conviction.",
                "tip": "💡 Memory Hack: Actor + Action + Object = Punchy Active Voice."
            },
            {
                "q": "Select the executive upgrade for 'very big impact':",
                "options": ["substantial impact", "huge big impact"],
                "correct": "substantial impact",
                "explanation": "Avoid intensifiers like 'very' or 'huge'. Use precise descriptors like 'substantial' or 'paramount'.",
                "tip": "💡 Memory Hack: Cut 'very' and choose a stronger adjective!"
            }
        ]
    },
    {
        "level_id": 5,
        "title": "Level 5: Master Linguistic Grandmaster (Boss Battle)",
        "description": "Complex relative clauses, subjunctive mood, and advanced syntax.",
        "icon": "⚡",
        "badge": "Linguistic Legend 🌟",
        "questions": [
            {
                "q": "The board requested that the CEO _____ present at the emergency briefing.",
                "options": ["be", "is"],
                "correct": "be",
                "explanation": "The subjunctive mood after demand/request verbs ('requested that...') uses the base verb form 'be'.",
                "tip": "💡 Memory Hack: Mandatory/Request verbs trigger base subjunctive: 'I insist that he BE present'."
            },
            {
                "q": "The architect _____ designed the zero-emission headquarters received the global award.",
                "options": ["who", "whom"],
                "correct": "who",
                "explanation": "'Who' functions as the subjective pronoun (doing the action). 'Whom' functions as an objective pronoun.",
                "tip": "💡 Memory Hack: Substitute 'he' or 'him'. 'He designed it' ➔ Who. 'Awarded to him' ➔ Whom."
            },
            {
                "q": "Neither the project manager nor the developers _____ able to replicate the timeout.",
                "options": ["were", "was"],
                "correct": "were",
                "explanation": "When subjects are joined by 'neither...nor', the verb agrees with the closer subject ('developers' = plural).",
                "tip": "💡 Memory Hack: Closer subject wins the agreement duel!"
            },
            {
                "q": "The committee reached a consensus, _____ pleased the executive sponsors.",
                "options": ["which", "that"],
                "correct": "which",
                "explanation": "Non-restrictive clauses following a comma that modify an entire previous statement use 'which', not 'that'.",
                "tip": "💡 Memory Hack: Comma + Which goes together; That hates commas!"
            },
            {
                "q": "If the architecture _____ optimized earlier, latency would be negligible now.",
                "options": ["had been", "was"],
                "correct": "had been",
                "explanation": "Counterfactual past conditions take past perfect subjunctive 'had been'.",
                "tip": "💡 Memory Hack: Imagined past scenarios require 'had been'."
            }
        ]
    }
]

DUOLINGO_BADGES = [
    {"id": "first_blood", "name": "Grammar Novice", "icon": "🌱", "desc": "Answered your first challenge correctly!"},
    {"id": "streak_3", "name": "Comet Streak", "icon": "🔥", "desc": "Achieved a 3-question streak without a mistake!"},
    {"id": "level_1", "name": "Agreement Apprentice", "icon": "🎖️", "desc": "Conquered Level 1: Subject-Verb Agreement!"},
    {"id": "level_2", "name": "Homophone Hunter", "icon": "🎯", "desc": "Conquered Level 2: Tricky Homophones!"},
    {"id": "level_3", "name": "Time Bender", "icon": "⏳", "desc": "Conquered Level 3: Modal & Tense Precision!"},
    {"id": "level_4", "name": "Executive Maestro", "icon": "👑", "desc": "Conquered Level 4: Executive Tone & Power Words!"},
    {"id": "level_5", "name": "Linguistic Legend", "icon": "🌟", "desc": "Completed all 5 Levels & Cleared the Duolingo Quest!"}
]

QUIZ_QUESTIONS = DUOLINGO_LEVELS[0]["questions"]

# ---------------------------------------------------------------------------
# SESSION STATE SETUP
# ---------------------------------------------------------------------------
if "user_text" not in st.session_state:
    st.session_state.user_text = ""
if "canvas_text" not in st.session_state:
    st.session_state.canvas_text = ""
if "whitelist" not in st.session_state:
    st.session_state.whitelist = set()
if "audit_results" not in st.session_state:
    st.session_state.audit_results = None
if "selected_dialect" not in st.session_state:
    st.session_state.selected_dialect = "en-US"
if "target_tone" not in st.session_state:
    st.session_state.target_tone = "Executive / Business 👔"
if "active_challenge" not in st.session_state:
    st.session_state.active_challenge = None
if "app_mode" not in st.session_state:
    st.session_state.app_mode = "studio"

# Duolingo Gamification State
if "duo_xp" not in st.session_state:
    st.session_state.duo_xp = 0
if "duo_hearts" not in st.session_state:
    st.session_state.duo_hearts = 5
if "duo_streak" not in st.session_state:
    st.session_state.duo_streak = 0
if "duo_unlocked_level" not in st.session_state:
    st.session_state.duo_unlocked_level = 1
if "duo_current_level" not in st.session_state:
    st.session_state.duo_current_level = 1
if "duo_q_idx" not in st.session_state:
    st.session_state.duo_q_idx = 0
if "duo_selected_answer" not in st.session_state:
    st.session_state.duo_selected_answer = None
if "duo_answer_evaluated" not in st.session_state:
    st.session_state.duo_answer_evaluated = False
if "duo_level_completed" not in st.session_state:
    st.session_state.duo_level_completed = False
if "duo_badges" not in st.session_state:
    st.session_state.duo_badges = set()

def perform_full_audit(text: str, dialect: str, target_tone: str, whitelist: set) -> Dict[str, Any]:
    if not text.strip():
        return None
    api_error = None
    raw_matches = []
    try:
        engine = get_cloud_tool(dialect)
        raw_matches = engine.check(text)
    except Exception as e:
        api_error = str(e)

    if api_error:
        return {"error": api_error}

    all_raw = check_custom_rules(text, raw_matches)

    matches = []
    for m in all_raw:
        tok = text[m.offset : m.offset + m.error_length]
        if not is_whitelisted(tok, whitelist):
            matches.append(m)

    corrected = language_tool_python.utils.correct(text, matches)

    diagnostics = []
    for m in matches:
        tok = text[m.offset : m.offset + m.error_length]
        diag = analyze_rule(m, tok)
        diag["token"] = tok
        diag["offset"] = m.offset
        diag["length"] = m.error_length
        diagnostics.append(diag)

    readability = calculate_readability(text)
    tone_data = analyze_tone_and_vocab(text, target_tone)

    return {
        "original": text,
        "corrected": corrected,
        "matches": matches,
        "raw_matches": all_raw,
        "diagnostics": diagnostics,
        "readability": readability,
        "tone": tone_data,
        "dialect": dialect,
        "target_tone": target_tone
    }

def re_filter_audit_results():
    if not st.session_state.audit_results or "raw_matches" not in st.session_state.audit_results:
        return
    res = st.session_state.audit_results
    text = res["original"]
    whitelist = st.session_state.whitelist

    matches = []
    for m in res["raw_matches"]:
        tok = text[m.offset : m.offset + m.error_length]
        if not is_whitelisted(tok, whitelist):
            matches.append(m)

    corrected = language_tool_python.utils.correct(text, matches)

    diagnostics = []
    for m in matches:
        tok = text[m.offset : m.offset + m.error_length]
        diag = analyze_rule(m, tok)
        diag["token"] = tok
        diag["offset"] = m.offset
        diag["length"] = m.error_length
        diagnostics.append(diag)

    res["matches"] = matches
    res["corrected"] = corrected
    res["diagnostics"] = diagnostics

def load_challenge(title: str, text: str, target_tone: str):
    st.session_state.canvas_text = text
    st.session_state.user_text = text
    st.session_state.target_tone = target_tone
    st.session_state.active_challenge = title
    cur_dialect = st.session_state.get("selected_dialect", "en-US")
    st.session_state.audit_results = perform_full_audit(text, cur_dialect, target_tone, st.session_state.whitelist)
    st.rerun()

# Duolingo Engine Handlers
def process_duo_answer(chosen: str, q_data: Dict[str, Any]):
    st.session_state.duo_selected_answer = chosen
    st.session_state.duo_answer_evaluated = True
    
    if chosen == q_data["correct"]:
        st.session_state.duo_xp += 20
        st.session_state.duo_streak += 1
        st.session_state.duo_badges.add("first_blood")
        if st.session_state.duo_streak >= 3:
            st.session_state.duo_badges.add("streak_3")
        st.balloons()
    else:
        st.session_state.duo_hearts = max(0, st.session_state.duo_hearts - 1)
        st.session_state.duo_streak = 0
    st.rerun()

def advance_duo_question():
    current_lvl = st.session_state.duo_current_level
    total_q = len(DUOLINGO_LEVELS[current_lvl - 1]["questions"])
    
    if st.session_state.duo_q_idx < total_q - 1:
        st.session_state.duo_q_idx += 1
        st.session_state.duo_selected_answer = None
        st.session_state.duo_answer_evaluated = False
        st.rerun()
    else:
        st.session_state.duo_level_completed = True
        st.session_state.duo_xp += 50
        st.session_state.duo_badges.add(f"level_{current_lvl}")
        if current_lvl == st.session_state.duo_unlocked_level and current_lvl < 5:
            st.session_state.duo_unlocked_level += 1
        if current_lvl == 5:
            st.session_state.duo_badges.add("level_5")
        st.rerun()

def render_trophy_cabinet():
    st.markdown("#### 🏆 **Achievement Rewards & Trophies**")
    b_cols = st.columns(len(DUOLINGO_BADGES))
    for i, b in enumerate(DUOLINGO_BADGES):
        is_owned = b["id"] in st.session_state.duo_badges
        card_class = "unlocked" if is_owned else ""
        icon = b["icon"] if is_owned else "🔒"
        status_txt = "UNLOCKED ✨" if is_owned else "Locked"
        with b_cols[i]:
            st.markdown(f"""
            <div class="badge-item {card_class}">
                <div class="badge-icon">{icon}</div>
                <div class="badge-name">{b['name']}</div>
                <div class="badge-desc">{b['desc']}</div>
                <div style="font-size: 0.68rem; font-weight: 800; color: {'#059669' if is_owned else '#94a3b8'}; margin-top: 6px;">{status_txt}</div>
            </div>
            """, unsafe_allow_html=True)

def render_language_arena():
    current_lvl_data = DUOLINGO_LEVELS[st.session_state.duo_current_level - 1]
    
    # 1. Top Gamification HUD
    hearts_display = "❤️ " * st.session_state.duo_hearts + "🖤 " * (5 - st.session_state.duo_hearts)
    
    hud_c1, hud_c2, hud_c3, hud_c4 = st.columns([2.5, 2.5, 2.5, 2.5])
    with hud_c1:
        st.markdown(f"""
        <div style="background: #ffffff; border: 1.5px solid #cbd5e1; border-radius: 14px; padding: 10px 16px; text-align: center; box-shadow: 0 4px 12px rgba(0,0,0,0.03);">
            <div style="font-size: 0.72rem; color: #64748b; font-weight: 700; text-transform: uppercase;">Hearts Remaining</div>
            <div style="font-size: 1.15rem; font-weight: 800;">{hearts_display}</div>
        </div>
        """, unsafe_allow_html=True)
    with hud_c2:
        st.markdown(f"""
        <div style="background: #ffffff; border: 1.5px solid #cbd5e1; border-radius: 14px; padding: 10px 16px; text-align: center; box-shadow: 0 4px 12px rgba(0,0,0,0.03);">
            <div style="font-size: 0.72rem; color: #64748b; font-weight: 700; text-transform: uppercase;">Combo Streak</div>
            <div style="font-size: 1.15rem; font-weight: 800; color: #ea580c;">🔥 {st.session_state.duo_streak} In A Row</div>
        </div>
        """, unsafe_allow_html=True)
    with hud_c3:
        st.markdown(f"""
        <div style="background: #ffffff; border: 1.5px solid #cbd5e1; border-radius: 14px; padding: 10px 16px; text-align: center; box-shadow: 0 4px 12px rgba(0,0,0,0.03);">
            <div style="font-size: 0.72rem; color: #64748b; font-weight: 700; text-transform: uppercase;">Total Experience</div>
            <div style="font-size: 1.15rem; font-weight: 800; color: #059669;">💎 {st.session_state.duo_xp} XP</div>
        </div>
        """, unsafe_allow_html=True)
    with hud_c4:
        st.markdown(f"""
        <div style="background: #ffffff; border: 1.5px solid #cbd5e1; border-radius: 14px; padding: 10px 16px; text-align: center; box-shadow: 0 4px 12px rgba(0,0,0,0.03);">
            <div style="font-size: 0.72rem; color: #64748b; font-weight: 700; text-transform: uppercase;">Trophy Cabinet</div>
            <div style="font-size: 1.15rem; font-weight: 800; color: #4338ca;">🏆 {len(st.session_state.duo_badges)} / {len(DUOLINGO_BADGES)} Unlocked</div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("<div style='height: 14px;'></div>", unsafe_allow_html=True)

    # 2. Grammar World Map (5 Levels)
    st.markdown("#### 🗺️ **Grammar World Map: Choose Your Level**")
    lvl_cols = st.columns(5)
    for idx, lvl in enumerate(DUOLINGO_LEVELS):
        l_num = lvl["level_id"]
        is_active = (l_num == st.session_state.duo_current_level)
        is_unlocked = (l_num <= st.session_state.duo_unlocked_level)
        is_cleared = (l_num < st.session_state.duo_unlocked_level)

        status_icon = "⭐ Cleared" if is_cleared else ("🎯 Active" if is_active else ("🔓 Open" if is_unlocked else "🔒 Locked"))
        btn_label = f"{lvl['icon']} Lvl {l_num}: {status_icon}"
        
        with lvl_cols[idx]:
            if st.button(btn_label, key=f"lvl_btn_{l_num}", use_container_width=True, disabled=not is_unlocked):
                st.session_state.duo_current_level = l_num
                st.session_state.duo_q_idx = 0
                st.session_state.duo_selected_answer = None
                st.session_state.duo_answer_evaluated = False
                st.session_state.duo_level_completed = False
                st.rerun()

    st.markdown("---")

    # 3. Heartbreak Screen
    if st.session_state.duo_hearts <= 0:
        st.markdown('<div class="duo-arena-card" style="text-align: center; background: #fff1f2; border-color: #fecdd3;">', unsafe_allow_html=True)
        st.markdown("### 💔 **You ran out of hearts!**")
        st.write("Mistakes are the stepping stones of grammar mastery! Take a quick breath and refill your hearts to jump back into the quest.")
        if st.button("💊 Refill 5 Hearts & Continue Training", type="primary"):
            st.session_state.duo_hearts = 5
            st.session_state.duo_answer_evaluated = False
            st.session_state.duo_selected_answer = None
            st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)
        render_trophy_cabinet()
        return

    # 4. Level Cleared Celebration Screen
    if st.session_state.duo_level_completed:
        st.balloons()
        st.markdown('<div class="duo-arena-card" style="text-align: center; background: linear-gradient(135deg, #f0fdf4 0%, #ecfdf5 100%); border-color: #86efac;">', unsafe_allow_html=True)
        st.markdown(f"## 🏆 **{current_lvl_data['title'].upper()} CLEARED!**")
        st.markdown(f"<p style='font-size: 1.15rem; color: #166534; font-weight: 700;'>Outstanding fluency! You mastered {current_lvl_data['description']}</p>", unsafe_allow_html=True)
        st.markdown(f"### 💎 **+50 Bonus XP Awarded!** | 🎖️ Unlocked: **{current_lvl_data['badge']}**")
        
        c_next1, c_next2, c_sp = st.columns([2.5, 2, 1.5])
        with c_next1:
            if st.session_state.duo_current_level < 5:
                if st.button("🚀 Unlock & Start Next Level ➔", type="primary", use_container_width=True):
                    st.session_state.duo_current_level += 1
                    st.session_state.duo_q_idx = 0
                    st.session_state.duo_selected_answer = None
                    st.session_state.duo_answer_evaluated = False
                    st.session_state.duo_level_completed = False
                    st.session_state.duo_hearts = 5
                    st.rerun()
            else:
                st.success("👑 YOU CONQUERED ALL 5 LEVELS! You are a Linguistic Legend!")
        with c_next2:
            if st.button("🔄 Replay This Level", use_container_width=True):
                st.session_state.duo_q_idx = 0
                st.session_state.duo_selected_answer = None
                st.session_state.duo_answer_evaluated = False
                st.session_state.duo_level_completed = False
                st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)
        render_trophy_cabinet()
        return

    # 5. Active Question Arena
    questions = current_lvl_data["questions"]
    q_idx = st.session_state.duo_q_idx
    q_data = questions[q_idx]

    # Progress Bar
    progress_val = (q_idx + 1) / len(questions)
    st.progress(progress_val)
    st.caption(f"**{current_lvl_data['title']}** • Challenge {q_idx + 1} of {len(questions)}")

    st.markdown('<div class="duo-arena-card">', unsafe_allow_html=True)
    
    # Mascot Speech Bubble
    st.markdown(f"""
    <div class="duo-mascot-row">
        <div class="duo-mascot-avatar">🦉</div>
        <div class="duo-bubble">
            "{q_data['q']}"
        </div>
    </div>
    """, unsafe_allow_html=True)

    # Answer Choices
    if not st.session_state.duo_answer_evaluated:
        col_a, col_b = st.columns(2)
        with col_a:
            if st.button(f"Option A: {q_data['options'][0]}", use_container_width=True, key=f"q_{st.session_state.duo_current_level}_{q_idx}_a"):
                process_duo_answer(q_data['options'][0], q_data)
        with col_b:
            if st.button(f"Option B: {q_data['options'][1]}", use_container_width=True, key=f"q_{st.session_state.duo_current_level}_{q_idx}_b"):
                process_duo_answer(q_data['options'][1], q_data)
    else:
        sel = st.session_state.duo_selected_answer
        is_corr = (sel == q_data["correct"])

        if is_corr:
            st.markdown(f"""
            <div class="duo-feedback-correct">
                <div style="font-size: 1.25rem; font-weight: 800; color: #15803d; margin-bottom: 6px;">
                    🎉 Spectacular! Correct Answer: "{q_data['correct']}" (+20 XP)
                </div>
                <div style="font-size: 0.95rem; line-height: 1.5; color: #166534;">
                    <b>Grammar Insight:</b> {q_data['explanation']}
                </div>
                <div style="font-size: 0.9rem; margin-top: 6px; color: #047857; font-weight: 600;">
                    {q_data['tip']}
                </div>
            </div>
            """, unsafe_allow_html=True)
        else:
            st.markdown(f"""
            <div class="duo-feedback-wrong">
                <div style="font-size: 1.25rem; font-weight: 800; color: #be123c; margin-bottom: 6px;">
                    💔 Not quite! (Lost 1 Heart)
                </div>
                <div style="font-size: 0.95rem; line-height: 1.5; color: #9f1239;">
                    <b>Correct Answer:</b> <u>{q_data['correct']}</u><br>
                    <b>Why:</b> {q_data['explanation']}
                </div>
                <div style="font-size: 0.9rem; margin-top: 6px; color: #881337; font-weight: 600;">
                    {q_data['tip']}
                </div>
            </div>
            """, unsafe_allow_html=True)

        st.markdown("<div style='height: 14px;'></div>", unsafe_allow_html=True)
        if st.button("Continue to Next Challenge ➔", type="primary"):
            advance_duo_question()

    st.markdown('</div>', unsafe_allow_html=True)

    # 6. Badges Cabinet
    render_trophy_cabinet()


# ---------------------------------------------------------------------------
# HERO NAVBAR WITH LOGO
# ---------------------------------------------------------------------------
logo_b64 = ""
if os.path.exists("logo.png"):
    try:
        with open("logo.png", "rb") as f_img:
            logo_b64 = base64.b64encode(f_img.read()).decode("utf-8")
    except Exception:
        logo_b64 = ""

logo_img_tag = f'<img src="data:image/png;base64,{logo_b64}" class="hero-logo-img" alt="Logo" />' if logo_b64 else '<div style="font-size: 2rem;">✨</div>'

st.markdown(f"""
<div class="hero-navbar">
    <div class="hero-brand">
        {logo_img_tag}
        <div>
            <h1 class="hero-title">Smart Grammar Checker</h1>
            <div class="hero-tagline">AI Grammar, Agreement & Syntax Mastery Studio • Live Online Engine</div>
        </div>
    </div>
    <div class="hero-pills">
        <div class="pill-online">
            <span class="pill-dot"></span>
            <span>Online Cloud Mode</span>
        </div>
    </div>
</div>
""", unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# APP NAVIGATION MODE: STUDIO vs. DUOLINGO QUEST ARENA
# ---------------------------------------------------------------------------
nav_c1, nav_c2 = st.columns([1, 1])
with nav_c1:
    is_studio = (st.session_state.app_mode == "studio")
    if st.button("✍️ AI Grammar Studio (Canvas, Files & Tone)", type="primary" if is_studio else "secondary", use_container_width=True):
        st.session_state.app_mode = "studio"
        st.rerun()

with nav_c2:
    is_duo = (st.session_state.app_mode == "duolingo")
    duo_title = f"🦉 Duolingo Grammar Arena (Lvl {st.session_state.duo_current_level} • {st.session_state.duo_xp} XP • {st.session_state.duo_hearts} ❤️)"
    if st.button(duo_title, type="primary" if is_duo else "secondary", use_container_width=True):
        st.session_state.app_mode = "duolingo"
        st.rerun()

st.markdown("<div style='height: 12px;'></div>", unsafe_allow_html=True)

# If in dedicated Duolingo Mode, render Arena and stop
if st.session_state.app_mode == "duolingo":
    render_language_arena()
    st.markdown("""
    <div class="app-footer">
        Smart Grammar Checker • Duolingo Grammar Quest Arena • Powered by Online Cloud API
    </div>
    """, unsafe_allow_html=True)
    st.stop()


# ---------------------------------------------------------------------------
# INTERACTIVE LEARNING PRESETS (1-CLICK LOAD)
# ---------------------------------------------------------------------------
st.markdown("##### 🚀 **Try Interactive Grammar Challenges (1-Click Learning):**")
if st.session_state.active_challenge:
    st.info(f"✨ Currently active challenge: **{st.session_state.active_challenge}** — 1-Click audit loaded. Detailed grammatical breakdown & coaching displayed below.")

c_p1, c_p2, c_p3, c_p4 = st.columns(4)

with c_p1:
    if st.button("📌 Subject-Verb Conflict", use_container_width=True):
        load_challenge(
            "Subject-Verb Conflict",
            "The group of researchers have concluded that each of the participants were affected by the noise.",
            "Academic / Scholarly 🎓"
        )
    st.caption("Tests collective noun ('have' ➔ 'has') & quantifier agreement ('were' ➔ 'was').")

with c_p2:
    if st.button("🔀 Tricky Homophones", use_container_width=True):
        load_challenge(
            "Tricky Homophones",
            "Their is a few mistake in the proposal and it's effect will be felt everyday by our team.",
            "Professional & Balanced 🌟"
        )
    st.caption("Tests there/their, few mistakes, its/it's, and everyday.")

with c_p3:
    if st.button("⏱️ Modal Auxiliary Slip", use_container_width=True):
        load_challenge(
            "Modal Auxiliary Slip",
            "I should of submitted the draft earlier because they was expecting it before noon.",
            "Casual / Friendly ☕"
        )
    st.caption("Tests 'should have' (never 'should of') and past tense agreement.")

with c_p4:
    if st.button("👔 Executive Business", use_container_width=True):
        load_challenge(
            "Executive Business",
            "Hey team, we needs to make sure that the spreadsheet is very big and we don't have no errors.",
            "Executive / Business 👔"
        )
    st.caption("Tests executive tone alignment, double negatives, and power vocabulary.")

st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)

# Dialect and Whitelist Toolstrip
ts_c1, ts_c2, ts_c3 = st.columns([3, 3, 3])
with ts_c1:
    st.caption("🌐 Regional Dialect")
    dialect = st.selectbox(
        "Regional Dialect",
        options=["en-US", "en-GB", "en-CA", "en-AU", "en-NZ"],
        format_func=lambda x: {
            "en-US": "🌐 US • English (United States)",
            "en-GB": "🇬🇧 UK • English (United Kingdom)",
            "en-CA": "🇨🇦 CA • English (Canada)",
            "en-AU": "🇦🇺 AU • English (Australia)",
            "en-NZ": "🇳🇿 NZ • English (New Zealand)"
        }.get(x, x),
        key="selected_dialect",
        label_visibility="collapsed"
    )
with ts_c2:
    st.caption("🎯 Target Voice & Tone")
    target_tone = st.selectbox(
        "Target Voice",
        options=["Executive / Business 👔", "Academic / Scholarly 🎓", "Direct & Impactful ⚡", "Casual / Friendly ☕"],
        key="target_tone",
        label_visibility="collapsed"
    )
with ts_c3:
    wl_count = len(st.session_state.whitelist)
    st.caption(f"🏷️ Technical Whitelist ({wl_count})")
    with st.popover(f"🏷️ Jargon Whitelist ({wl_count} terms)", use_container_width=True):
        st.markdown("**Domain Jargon & Whitelisted Acronyms**")
        st.caption("Whitelisted terms are exempt from spelling and grammar flags during audits.")

        def handle_add_whitelist():
            w = st.session_state.get("new_wl_input", "").strip().lower()
            if w:
                st.session_state.whitelist.add(w)
                st.session_state["new_wl_input"] = ""
                re_filter_audit_results()

        st.text_input(
            "Add custom jargon",
            placeholder="e.g. Kubernetes, PyTorch, Sem...",
            key="new_wl_input",
            on_change=handle_add_whitelist,
            label_visibility="collapsed"
        )
        if st.session_state.whitelist:
            st.markdown("<div style='height: 6px;'></div>", unsafe_allow_html=True)
            st.write("**Active Whitelisted Terms:**")
            for term in sorted(list(st.session_state.whitelist)):
                w_col1, w_col2 = st.columns([4, 1])
                with w_col1:
                    st.code(term, language="text")
                with w_col2:
                    if st.button("✖", key=f"del_wl_{term}", help=f"Remove '{term}' from whitelist"):
                        st.session_state.whitelist.discard(term)
                        re_filter_audit_results()
                        st.rerun()
            if st.button("🧹 Clear All Whitelist", use_container_width=True):
                st.session_state.whitelist.clear()
                re_filter_audit_results()
                st.rerun()
        else:
            st.caption("No custom jargon added yet. Type a term above and press Enter.")

# Studio Navigation Tabs
tab_editor, tab_files, tab_tone, tab_quiz, tab_guide = st.tabs([
    "✍️ AI Grammar Canvas",
    "📂 Document & File Audit",
    "⚡ Tone & Vocabulary Booster",
    "🎮 Grammar Quiz & Game",
    "📚 Master Linguistic Rules"
])

# ===========================================================================
# TAB 1: AI GRAMMAR CANVAS
# ===========================================================================
with tab_editor:
    editor_input = st.text_area(
        "Canvas",
        key="canvas_text",
        height=190,
        placeholder="Type or paste any paragraph, proposal, essay, or email to inspect grammar, agreement, spelling, and tone...",
        label_visibility="collapsed"
    )
    st.session_state.user_text = editor_input

    btn_c1, btn_c2, btn_c3, btn_sp = st.columns([1.8, 1.4, 1.8, 3.2])
    with btn_c1:
        run_audit = st.button("⚡ Inspect & Learn", type="primary", use_container_width=True)
    with btn_c2:
        if st.button("🧹 Clear Canvas", use_container_width=True):
            st.session_state.canvas_text = ""
            st.session_state.user_text = ""
            st.session_state.audit_results = None
            st.session_state.active_challenge = None
            st.rerun()
    with btn_c3:
        if st.session_state.audit_results and st.button("✨ Auto-Fix All", use_container_width=True):
            corr_text = st.session_state.audit_results["corrected"]
            st.session_state.canvas_text = corr_text
            st.session_state.user_text = corr_text
            st.session_state.audit_results = None
            st.success("All grammar fixes applied instantly!")
            st.rerun()

    # Grammar Engine Execution
    if run_audit:
        if not editor_input.strip():
            st.warning("⚠️ Please provide text to inspect.")
        else:
            with st.spinner("Analyzing grammar, agreement, and syntax via Online Cloud API..."):
                audit_res = perform_full_audit(editor_input, dialect, target_tone, st.session_state.whitelist)
                if audit_res and "error" in audit_res:
                    st.error(f"🌐 Cloud API connection issue: {audit_res['error']}. Please check your connection or retry in a moment.")
                else:
                    st.session_state.audit_results = audit_res
                    st.session_state.active_challenge = None
                    if audit_res and len(audit_res["diagnostics"]) == 0:
                        st.balloons()

    # Results Dashboard
    if st.session_state.audit_results:
        res = st.session_state.audit_results
        orig = res["original"]
        corr = res["corrected"]
        diag = res["diagnostics"]
        stats = res["readability"]
        tone = res["tone"]

        # Check if target tone was updated while viewing results
        if res.get("target_tone") != target_tone:
            res["target_tone"] = target_tone
            res["tone"] = analyze_tone_and_vocab(res["original"], target_tone)
            tone = res["tone"]

        # Check if dialect was changed
        if res.get("dialect") != dialect:
            st.info(f"🌐 Regional dialect changed from **{res.get('dialect')}** to **{dialect}**. Click '⚡ Inspect & Learn' to re-audit with {dialect} spelling conventions.")

        # 1. Gamified Metrics Ribbon
        quality_score = max(0, 100 - (len(diag) * 4))
        q_color = "tile-val-emerald" if quality_score >= 85 else "tile-val-amber" if quality_score >= 60 else "tile-val-rose"
        target_score = tone.get("target_alignment_score", 50)
        t_color = "tile-val-emerald" if target_score >= 75 else "tile-val-amber" if target_score >= 50 else "tile-val-rose"

        st.markdown(f"""
        <div class="metrics-ribbon">
            <div class="metric-tile">
                <div class="tile-val {q_color}">{quality_score}%</div>
                <div class="tile-lbl">Mastery Score</div>
            </div>
            <div class="metric-tile">
                <div class="tile-val tile-val-amber">{len(diag)}</div>
                <div class="tile-lbl">Issues Found</div>
            </div>
            <div class="metric-tile">
                <div class="tile-val tile-val-blue">{stats['words']}</div>
                <div class="tile-lbl">Word Count</div>
            </div>
            <div class="metric-tile">
                <div class="tile-val tile-val-purple">Gr. {stats['grade']}</div>
                <div class="tile-lbl">Reading Level</div>
            </div>
            <div class="metric-tile">
                <div class="tile-val tile-val-rose">{tone['persona'].split()[0]}</div>
                <div class="tile-lbl">Detected Tone</div>
            </div>
            <div class="metric-tile">
                <div class="tile-val {t_color}">{target_score}%</div>
                <div class="tile-lbl">Target Match</div>
            </div>
        </div>
        """, unsafe_allow_html=True)

        # Target Tone Strategic Alignment Card
        target_status = tone.get("target_alignment_status", "Ready")
        coaching_items = tone.get("target_coaching", [])
        status_color = "#059669" if target_score >= 75 else "#d97706" if target_score >= 50 else "#e11d48"

        st.markdown(f"""
        <div style="background: linear-gradient(135deg, #f8fafc 0%, #eff6ff 100%); border: 1.5px solid #bfdbfe; border-radius: 12px; padding: 14px 18px; margin: 12px 0 18px 0; box-shadow: 0 2px 8px rgba(59,130,246,0.06);">
            <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 8px;">
                <div style="font-weight: 800; color: #1e3a8a; font-size: 1.02rem;">🎯 Target Voice: {tone.get('target_tone', target_tone)}</div>
                <div style="font-size: 0.88rem; font-weight: 700; color: {status_color};">
                    {target_status} • {target_score}% Match
                </div>
            </div>
            <div style="font-size: 0.9rem; color: #334155; margin-bottom: 4px;">
                <b>Detected Persona:</b> {tone['persona']} — <i>{tone['summary']}</i>
            </div>
        </div>
        """, unsafe_allow_html=True)

        if coaching_items:
            with st.expander(f"💡 Actionable Voice Recommendations for {tone.get('target_tone', target_tone)}", expanded=(target_score < 75)):
                for tip in coaching_items:
                    st.markdown(f"- {tip}")

        # 2. Side-by-Side Comparison
        col_orig, col_corr = st.columns(2)
        with col_orig:
            st.markdown('<div class="glow-card">', unsafe_allow_html=True)
            st.markdown("#### 📌 Original Draft (Annotated Mistakes)")
            annotated_html = build_annotated_preview(orig, res["matches"])
            st.markdown(f'<div style="line-height: 1.85; font-size: 1.02rem; padding: 6px 0;">{annotated_html}</div>', unsafe_allow_html=True)
            st.markdown('</div>', unsafe_allow_html=True)

        with col_corr:
            st.markdown('<div class="glow-card">', unsafe_allow_html=True)
            st.markdown("#### ✨ Tuned & Polished Masterpiece")
            st.markdown(f'<div style="line-height: 1.85; font-size: 1.02rem; color: #065f46; font-weight: 600; padding: 6px 0;">{html.escape(corr)}</div>', unsafe_allow_html=True)
            st.code(corr, language="text")
            st.markdown('</div>', unsafe_allow_html=True)

        # 3. Exact Word Diff
        with st.expander("🔍 Inspect Exact Word-by-Word Diff", expanded=False):
            diff_markup = build_visual_diff(orig, corr)
            st.markdown(f'<div class="glow-card" style="line-height: 2.1;">{diff_markup}</div>', unsafe_allow_html=True)

        # 4. Interactive Learning Flashcards ("Why is this wrong?" & "Why is this right?")
        st.markdown('<div class="glow-card">', unsafe_allow_html=True)
        st.markdown("### 🎓 **AI Grammar Tutor: Master Your Mistakes**")
        st.caption("Learn from every error with detailed grammatical rules, explanations, and pro tips:")

        if not diag:
            st.success("🎉 Phenomenal writing! No grammar, agreement, or spelling errors identified.")
        else:
            # Category Filter
            available_cats = ["All Categories"] + sorted(list({d["category"] for d in diag}))
            filter_sel = st.radio("Filter issues by type:", options=available_cats, horizontal=True)

            filtered = [d for d in diag if filter_sel == "All Categories" or d["category"] == filter_sel]

            for d in filtered:
                st.markdown(f"""
                <div class="tutor-card">
                    <div class="tutor-header">
                        <div>
                            <span class="cat-tag {d['tag']}">{d['category']}</span>
                            <span class="token-compare" style="margin-left: 8px;">
                                <span class="tok-wrong">{html.escape(d['token'])}</span>
                                <span class="tok-arrow">➔</span>
                                <span class="tok-right">{html.escape(d['best_replacement'])}</span>
                            </span>
                        </div>
                        <span style="font-size: 0.76rem; color: #475569; font-weight: 600; font-family: 'JetBrains Mono', monospace;">Rule: {d['rule_id']}</span>
                    </div>
                    <div class="tutor-block tb-wrong">
                        <b>🔴 Why is this wrong?</b><br>
                        {d['why_wrong']}
                    </div>
                    <div class="tutor-block tb-right">
                        <b>🟢 Why is this right?</b><br>
                        {d['why_right']}
                    </div>
                    <div class="tutor-block tb-protip">
                        {d.get('protip', '💡 Always review subject-verb pairs to build natural fluency.')}
                    </div>
                </div>
                """, unsafe_allow_html=True)

        st.markdown('</div>', unsafe_allow_html=True)

        # 5. Export Actions
        st.markdown("### 💾 **Download & Share Polished Document**")
        exp1, exp2, exp3 = st.columns(3)
        with exp1:
            st.download_button(
                label="📥 Download Corrected Text (.txt)",
                data=corr,
                file_name="polished_text.txt",
                mime="text/plain",
                use_container_width=True
            )
        with exp2:
            docx_buffer = export_docx_file(orig, corr, diag, dialect)
            st.download_button(
                label="📄 Download Formatted Word (.docx)",
                data=docx_buffer,
                file_name="polished_document.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
        with exp3:
            audit_json = {
                "dialect": dialect,
                "engine": "Online Cloud API",
                "quality_score": quality_score,
                "readability": stats,
                "tone": tone,
                "total_issues": len(diag),
                "issues": diag
            }
            st.download_button(
                label="📊 Download Full Audit Report (.json)",
                data=json.dumps(audit_json, indent=2),
                file_name="grammar_audit_report.json",
                mime="application/json",
                use_container_width=True
            )

# ===========================================================================
# TAB 2: DOCUMENT & FILE AUDIT
# ===========================================================================
with tab_files:
    st.markdown('<div class="glow-card">', unsafe_allow_html=True)
    st.markdown("### 📂 **Multi-Format Document Audit**")
    st.write("Upload any document to audit grammar, spelling, agreement, and style in one unified pass.")
    st.caption("Supported formats: **Word (.docx, .doc)**, **PDF (.pdf)**, **Plain Text (.txt)**, **Markdown (.md)**, **CSV (.csv)**")

    uploaded = st.file_uploader(
        "Upload a document",
        type=["docx", "doc", "pdf", "txt", "md", "csv"],
        label_visibility="collapsed"
    )

    if uploaded is not None:
        file_text, f_meta, f_err = parse_uploaded_file(uploaded)
        if f_err:
            st.error(f"❌ {f_err}")
        elif not file_text.strip():
            st.warning("⚠️ The uploaded document contains no readable text.")
        else:
            st.info(f"📄 Loaded **{f_meta['filename']}** ({f_meta['size_kb']} KB, ~{f_meta['word_count']} words)")

            if st.button("🚀 Audit Entire Document with Cloud Engine", type="primary"):
                with st.spinner("Extracting paragraphs and auditing with Cloud Engine..."):
                    api_error = None
                    raw_doc_matches = []
                    try:
                        engine = get_cloud_tool(dialect)
                        raw_doc_matches = engine.check(file_text)
                    except Exception as e:
                        api_error = str(e)

                    if api_error:
                        st.error(f"🌐 Cloud API connection issue: {api_error}. Please check your connection or retry in a moment.")
                    else:
                        doc_matches = []
                        for m in raw_doc_matches:
                            tok = file_text[m.offset : m.offset + m.error_length]
                            if not is_whitelisted(tok, st.session_state.whitelist):
                                doc_matches.append(m)

                        doc_corrected = language_tool_python.utils.correct(file_text, doc_matches)
                        doc_stats = calculate_readability(file_text)
                        doc_tone = analyze_tone_and_vocab(file_text, target_tone)

                        doc_diagnostics = []
                        for m in doc_matches:
                            tok = file_text[m.offset : m.offset + m.error_length]
                            diag = analyze_rule(m, tok)
                            diag["token"] = tok
                            diag["offset"] = m.offset
                            diag["length"] = m.error_length
                            doc_diagnostics.append(diag)

                        st.success(f"Audit completed: {len(doc_diagnostics)} grammatical or stylistic improvements identified.")

                        # Metrics Strip
                        st.markdown(f"""
                        <div class="metrics-ribbon">
                            <div class="metric-tile">
                                <div class="tile-val tile-val-blue">{doc_stats['words']}</div>
                                <div class="tile-lbl">Total Words</div>
                            </div>
                            <div class="metric-tile">
                                <div class="tile-val tile-val-amber">{len(doc_diagnostics)}</div>
                                <div class="tile-lbl">Issues Resolved</div>
                            </div>
                            <div class="metric-tile">
                                <div class="tile-val tile-val-purple">Gr. {doc_stats['grade']}</div>
                                <div class="tile-lbl">Target Grade</div>
                            </div>
                            <div class="metric-tile">
                                <div class="tile-val tile-val-emerald">{doc_stats['level'].split()[0]}</div>
                                <div class="tile-lbl">Reading Level</div>
                            </div>
                            <div class="metric-tile">
                                <div class="tile-val tile-val-rose">{doc_tone['persona'].split()[0]}</div>
                                <div class="tile-lbl">Document Tone</div>
                            </div>
                        </div>
                        """, unsafe_allow_html=True)

                        # Previews
                        col_p1, col_p2 = st.columns(2)
                        with col_p1:
                            st.markdown("**Original Document Preview:**")
                            st.text_area("OrigPreview", value=file_text[:3000] + ("..." if len(file_text) > 3000 else ""), height=220, disabled=True, label_visibility="collapsed")
                        with col_p2:
                            st.markdown("**Tuned & Corrected Document Preview:**")
                            st.text_area("CorrPreview", value=doc_corrected[:3000] + ("..." if len(doc_corrected) > 3000 else ""), height=220, disabled=True, label_visibility="collapsed")

                        # Downloads
                        st.markdown("#### 📥 **Download Polished Document**")
                        f_d1, f_d2 = st.columns(2)
                        with f_d1:
                            out_doc = export_docx_file(file_text, doc_corrected, doc_diagnostics, dialect)
                            st.download_button(
                                label="📄 Download Corrected Word Document (.docx)",
                                data=out_doc,
                                file_name=f"polished_{uploaded.name.rsplit('.', 1)[0]}.docx",
                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                use_container_width=True
                            )
                        with f_d2:
                            st.download_button(
                                label="📝 Download Corrected Text (.txt)",
                                data=doc_corrected,
                                file_name=f"polished_{uploaded.name.rsplit('.', 1)[0]}.txt",
                                mime="text/plain",
                                use_container_width=True
                            )
    st.markdown('</div>', unsafe_allow_html=True)

# ===========================================================================
# TAB 3: TONE & VOCABULARY BOOSTER
# ===========================================================================
with tab_tone:
    st.markdown('<div class="glow-card">', unsafe_allow_html=True)
    st.markdown("### ⚡ **Tone Mastery & Vocabulary Power-Up**")
    st.write("Evaluate your writing voice and discover strong vocabulary alternatives for common plain words.")

    sample_for_tone = st.session_state.canvas_text if st.session_state.canvas_text else (st.session_state.user_text if st.session_state.user_text else "Please inspect some text in Tab 1.")
    active_tone = analyze_tone_and_vocab(sample_for_tone, target_tone)

    st.markdown(f"""
    <div style="background: linear-gradient(135deg, #ede9fe 0%, #f0fdf4 100%); border: 1px solid #c4b5fd; border-radius: 12px; padding: 18px 22px; margin-bottom: 16px;">
        <div style="font-size: 1.15rem; font-weight: 800; color: #1e1b4b;">Detected Writing Voice: {active_tone['persona']}</div>
        <div style="font-size: 0.9rem; color: #475569; margin-top: 4px;">{active_tone['summary']}</div>
    </div>
    """, unsafe_allow_html=True)

    t_score = active_tone.get("target_alignment_score", 50)
    t_status = active_tone.get("target_alignment_status", "Ready")
    t_badge_col = "#059669" if t_score >= 75 else "#d97706" if t_score >= 50 else "#e11d48"

    st.markdown(f"""
    <div style="background: #ffffff; border: 1.5px solid #dbeafe; border-radius: 12px; padding: 16px 20px; margin-bottom: 20px; box-shadow: 0 2px 8px rgba(0,0,0,0.02);">
        <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 8px;">
            <span style="font-weight: 800; color: #1e40af; font-size: 1.05rem;">🎯 Target Voice Alignment: {active_tone['target_tone']}</span>
            <span style="font-weight: 800; color: {t_badge_col}; font-size: 0.95rem;">
                {t_status} • {t_score}% Match
            </span>
        </div>
        <div style="margin-bottom: 12px;">
            <div style="background: #e2e8f0; border-radius: 9999px; height: 10px; overflow: hidden;">
                <div style="background: {'#10b981' if t_score >= 75 else '#f59e0b' if t_score >= 50 else '#f43f5e'}; width: {t_score}%; height: 100%;"></div>
            </div>
        </div>
        <div style="font-size: 0.9rem; font-weight: 700; color: #334155; margin-bottom: 4px;">
            Strategic Voice Elevation Recommendations:
        </div>
        <ul style="font-size: 0.88rem; color: #475569; margin-top: 4px; margin-bottom: 0;">
            {''.join(f'<li>{tip}</li>' for tip in active_tone['target_coaching'])}
        </ul>
    </div>
    """, unsafe_allow_html=True)

    tc1, tc2, tc3 = st.columns(3)
    with tc1:
        st.markdown(f"""
        <div class="tone-meter-bar">
            <div style="display: flex; justify-content: space-between; font-weight: 700; color: #1e40af; margin-bottom: 6px;">
                <span>Formality Index</span>
                <span>{active_tone['formality']}%</span>
            </div>
        </div>
        """, unsafe_allow_html=True)
        st.progress(active_tone['formality'] / 100)
        st.caption("High formality eliminates slang and contractions.")

    with tc2:
        st.markdown(f"""
        <div class="tone-meter-bar">
            <div style="display: flex; justify-content: space-between; font-weight: 700; color: #065f46; margin-bottom: 6px;">
                <span>Confidence & Directness</span>
                <span>{active_tone['directness']}%</span>
            </div>
        </div>
        """, unsafe_allow_html=True)
        st.progress(active_tone['directness'] / 100)
        st.caption("Active voice and zero tentative hedging phrases.")

    with tc3:
        st.markdown(f"""
        <div class="tone-meter-bar">
            <div style="display: flex; justify-content: space-between; font-weight: 700; color: #9d174d; margin-bottom: 6px;">
                <span>Diplomacy & Politeness</span>
                <span>{active_tone['politeness']}%</span>
            </div>
        </div>
        """, unsafe_allow_html=True)
        st.progress(active_tone['politeness'] / 100)
        st.caption("Courteous phrases and diplomatic framing.")

    # Vocabulary Power-Up Section
    st.markdown("---")
    st.markdown("#### 💎 **Vocabulary Power-Up: Upgrade Plain Words**")
    if active_tone["vocab_suggestions"]:
        st.write("We detected plain or overused phrases in your draft. Upgrade them with these strong alternatives:")
        for v in active_tone["vocab_suggestions"]:
            st.markdown(f"""
            <div style="background: #ffffff; border: 1px solid #e2e8f0; border-radius: 10px; padding: 10px 16px; margin-bottom: 8px; box-shadow: 0 2px 6px rgba(0,0,0,0.02);">
                <span style="color: #be123c; text-decoration: line-through; font-weight: 600;">{v['weak']}</span> 
                ➔ <span style="color: #15803d; font-weight: 700;">{v['strong']}</span>
            </div>
            """, unsafe_allow_html=True)
    else:
        st.info("💡 Tip: Use power verbs and rich adjectives to make your writing engaging! Try using phrases like 'very big', 'very good', or 'make sure' in the editor to see instant upgrade recommendations.")

    st.markdown('</div>', unsafe_allow_html=True)

# ===========================================================================
# TAB 4: GRAMMAR QUIZ & GAME
# ===========================================================================
with tab_quiz:
    st.markdown('<div class="quiz-box">', unsafe_allow_html=True)
    st.markdown("### 🦉 **Duolingo Grammar Arena: Learn Through Play**")
    st.write("Play interactive grammar levels with Hearts ❤️, Streaks 🔥, XP points 💎, and unlockable achievement trophies 🏆!")

    col_q1, col_q2 = st.columns([1.5, 2.5])
    with col_q1:
        st.markdown(f"""
        <div style="background: #ffffff; border: 1.5px solid #cbd5e1; border-radius: 14px; padding: 18px; box-shadow: 0 4px 12px rgba(0,0,0,0.03);">
            <div style="font-size: 1.1rem; font-weight: 800; color: #4338ca; margin-bottom: 8px;">🎮 Current Quest Status</div>
            <div style="margin-bottom: 6px;"><b>Stage:</b> Level {st.session_state.duo_current_level} of 5</div>
            <div style="margin-bottom: 6px;"><b>Hearts:</b> {'❤️ ' * st.session_state.duo_hearts}</div>
            <div style="margin-bottom: 6px;"><b>Streak:</b> 🔥 {st.session_state.duo_streak} In A Row</div>
            <div style="margin-bottom: 6px;"><b>Total XP:</b> 💎 {st.session_state.duo_xp} XP</div>
            <div><b>Badges:</b> 🏆 {len(st.session_state.duo_badges)} / {len(DUOLINGO_BADGES)} Unlocked</div>
        </div>
        """, unsafe_allow_html=True)
        st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)
        if st.button("🚀 Enter Full Duolingo Arena Mode ➔", type="primary", use_container_width=True):
            st.session_state.app_mode = "duolingo"
            st.rerun()

    with col_q2:
        if st.session_state.duo_hearts <= 0:
            st.warning("💔 You ran out of hearts in the Arena!")
            if st.button("💊 Refill 5 Hearts", key="tab_refill", type="primary"):
                st.session_state.duo_hearts = 5
                st.session_state.duo_answer_evaluated = False
                st.session_state.duo_selected_answer = None
                st.rerun()
        elif st.session_state.duo_level_completed:
            st.success(f"🎉 Level {st.session_state.duo_current_level} Cleared (+50 XP)!")
            if st.session_state.duo_current_level < 5:
                if st.button("🚀 Unlock & Start Next Level ➔", key="tab_next_lvl", type="primary"):
                    st.session_state.duo_current_level += 1
                    st.session_state.duo_q_idx = 0
                    st.session_state.duo_selected_answer = None
                    st.session_state.duo_answer_evaluated = False
                    st.session_state.duo_level_completed = False
                    st.session_state.duo_hearts = 5
                    st.rerun()
            else:
                st.info("👑 All 5 levels completed! You are a Linguistic Legend.")
                if st.button("🔄 Replay Level 1", key="tab_replay_lvl"):
                    st.session_state.duo_current_level = 1
                    st.session_state.duo_q_idx = 0
                    st.session_state.duo_selected_answer = None
                    st.session_state.duo_answer_evaluated = False
                    st.session_state.duo_level_completed = False
                    st.rerun()
        else:
            cur_lvl = DUOLINGO_LEVELS[st.session_state.duo_current_level - 1]
            q_idx = st.session_state.duo_q_idx
            q_data = cur_lvl["questions"][q_idx]

            st.markdown(f"""
            <div style="font-size: 1.05rem; font-weight: 700; margin-bottom: 12px; color: #0f172a;">
                <b>{cur_lvl['title']}</b> (Question {q_idx + 1} of {len(cur_lvl['questions'])})<br>
                <span style="color: #0369a1; font-size: 1.15rem;">"{q_data['q']}"</span>
            </div>
            """, unsafe_allow_html=True)

            if not st.session_state.duo_answer_evaluated:
                c1, c2 = st.columns(2)
                with c1:
                    if st.button(f"Option A: {q_data['options'][0]}", use_container_width=True, key="tab_opt_a"):
                        process_duo_answer(q_data['options'][0], q_data)
                with c2:
                    if st.button(f"Option B: {q_data['options'][1]}", use_container_width=True, key="tab_opt_b"):
                        process_duo_answer(q_data['options'][1], q_data)
            else:
                sel = st.session_state.duo_selected_answer
                if sel == q_data["correct"]:
                    st.success(f"🎉 Correct! '{sel}' (+20 XP)")
                else:
                    st.error(f"❌ Not quite! Correct is '{q_data['correct']}' (-1 Heart)")
                st.info(f"💡 {q_data['explanation']}\n\n{q_data['tip']}")
                if st.button("Next Challenge ➔", key="tab_next"):
                    advance_duo_question()

    st.markdown('</div>', unsafe_allow_html=True)

# ===========================================================================
# TAB 5: MASTER LINGUISTIC RULES
# ===========================================================================
with tab_guide:
    st.markdown('<div class="glow-card">', unsafe_allow_html=True)
    st.markdown("### 📚 **Master Linguistic Rules Guide**")
    st.write("Understand the core principles behind standard English grammar, syntax, and style:")

    with st.expander("📌 1. Subject-Verb Agreement Principles", expanded=True):
        st.markdown("""
        - **Third-Person Singular Present Simple**: Singular subjects (*he, she, it, the company, the scientist*) require verbs ending in **-s** or **-es** (*She conducts, The system runs*).
        - **Distributives (*each / every*)**: Distributive pronouns govern singular verbs (*Each of the options is valid*, not *are*).
        - **Quantifiers (*a few / many*)**: Govern plural countable nouns (*A few revisions*, not *A few revision*).
        """)

    with st.expander("🔀 2. Frequently Confused Homophones", expanded=False):
        st.markdown("""
        - **There / Their / They're**:
          - *There*: Location or existential clause (*There is an issue*).
          - *Their*: Possessive determiner (*Their server is live*).
          - *They're*: Contraction for *they are* (*They're reviewing the code*).
        - **Its vs. It's**:
          - *Its*: Possessive form (*The model updated its weights*).
          - *It's*: Contraction of *it is* (*It's completed*).
        - **Affect vs. Effect**:
          - *Affect (Verb)*: To influence (*Weather affects performance*).
          - *Effect (Noun)*: The consequence (*The effect of optimization*).
        """)

    with st.expander("⏱️ 3. Modal Auxiliary Constructions", expanded=False):
        st.markdown("""
        - **Should have vs. Should of**: Modal auxiliary verbs (*should, could, would*) are paired with the auxiliary verb **have** (*should have* or *should've*). The phrasing *should of* is always grammatically incorrect.
        """)
    st.markdown('</div>', unsafe_allow_html=True)

# ---------------------------------------------------------------------------
# STUDIO FOOTER
# ---------------------------------------------------------------------------
st.markdown("""
<div class="app-footer">
    Smart Grammar Checker • AI-Powered Grammar, Syntax & Tone Mastery Studio • Powered by Online Cloud API
</div>
""", unsafe_allow_html=True)